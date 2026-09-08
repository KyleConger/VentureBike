"""Add the exact Q3 Advanced Balanced Scorecard (all 7 firms) to Q3Data.xlsx.

Supersedes the chart-read BSC block in Q3_Industry_Graphs (Bike Bros and
Spoke'd Up were mis-read). Source: in-sim Advanced Balanced Scorecard, Q3.
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DST = "quarters/Q3/Q3Data.xlsx"

# Column order matches the in-sim table.
FIRMS = [
    "Bike Bros", "LiteCycle", "WeBike", "BB LLC",
    "Spoke'd Up", "SpaceBikes", "MILC Bikes",
]

# indicator -> scores in FIRMS order
SCORES = {
    "Total Performance":          [0.900, 2.765, 0.411, 21.531, 2.316, 10.753, 2.604],
    "Financial Performance":      [12.856, 14.840, 5.319, 37.687, 14.047, 33.398, 14.895],
    "Market Performance":         [0.085, 0.153, 0.098, 0.345, 0.172, 0.270, 0.170],
    "Marketing Effectiveness":    [0.765, 0.758, 0.738, 0.765, 0.660, 0.763, 0.755],
    "Investment in Future":       [7.401, 5.657, 6.795, 3.741, 5.601, 3.466, 5.038],
    "Wealth":                     [0.797, 0.768, 0.668, 0.957, 0.803, 1.017, 0.806],
    "Human Resource Management":  [0.691, 0.732, 0.713, 0.767, 0.719, 0.735, 0.713],
    "Asset Management":           [0.366, 0.539, 0.353, 0.920, 0.479, 0.666, 0.473],
    "Manufacturing Productivity": [0.722, 0.938, 0.938, 0.856, 0.938, 0.906, 0.995],
    "Financial Risk":             [1.000, 1.000, 1.000, 1.000, 1.000, 1.000, 1.000],
}

COMPONENTS = [k for k in SCORES if k != "Total Performance"]
WEBIKE = FIRMS.index("WeBike")

wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
bad = PatternFill("solid", fgColor="FFC7CE")
good = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFEB9C")
ours = PatternFill("solid", fgColor="FFF2CC")
thin = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)


def ranks(values, higher_is_better=True):
    """Competition rank (1 = best). Ties share the best rank in the group."""
    indexed = list(enumerate(values))
    indexed.sort(key=lambda x: -x[1] if higher_is_better else x[1])
    out = [0] * len(values)
    i = 0
    while i < len(indexed):
        j = i
        while j < len(indexed) and indexed[j][1] == indexed[i][1]:
            j += 1
        for k in range(i, j):
            out[indexed[k][0]] = i + 1
        i = j
    return out


def block(ws, start, rows, wrap_col=None):
    for i, row in enumerate(rows):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(start + i, c, v)
            cell.border = thin
            if i == 0:
                cell.fill = hdr
                cell.font = bold
            elif wrap_col and c == wrap_col:
                cell.alignment = Alignment(wrap_text=True)
    return start + len(rows) + 2


if "Q3_Advanced_BSC" in wb.sheetnames:
    del wb["Q3_Advanced_BSC"]
ws = wb.create_sheet("Q3_Advanced_BSC", 1)

ws["A1"] = "Advanced Balanced Scorecard — Q3 actual (all 7 firms)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = (
    "Exact in-sim numbers. Total Performance is the product of the nine indicators. "
    "Supersedes the chart-read BSC block in Q3_Industry_Graphs (Bike Bros 0.900 not ~2.5; "
    "Spoke'd Up 2.316 not ~1.0)."
)
ws["A2"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A2:L2")
ws.row_dimensions[2].height = 36

# ---- matrix ----
r = 4
header = ["Indicator"] + FIRMS + ["Min", "Max", "Average", "WeBike rank"]
rows = [header]
webike_ranks = {}
for name, vals in SCORES.items():
    rk = ranks(vals)
    webike_ranks[name] = rk[WEBIKE]
    rows.append(
        [name] + list(vals) + [min(vals), max(vals), sum(vals) / 7, rk[WEBIKE]]
    )
r2 = block(ws, r, rows)
for i, (name, vals) in enumerate(SCORES.items()):
    row = r + 1 + i
    best = max(vals)
    rk = webike_ranks[name]
    for c, v in enumerate(vals, start=2):
        ws.cell(row, c).number_format = "0.000"
        if v == best:
            ws.cell(row, c).fill = good
    # WeBike column: rank color wins over "best" (only FR is both)
    if name == "Financial Risk" or rk <= 2:
        ws.cell(row, 4).fill = good
        ws.cell(row, 12).fill = good
    elif rk >= 6:
        ws.cell(row, 4).fill = bad
        ws.cell(row, 12).fill = bad
    else:
        ws.cell(row, 4).fill = ours
    ws.cell(row, 4).font = bold
    for c in range(9, 12):
        ws.cell(row, c).number_format = "0.000"
r = r2

# ---- WeBike rank card ----
rows = [
    ["WeBike rank card", "Score", "Rank / 7", "Best firm", "Best score", "Read"],
    ["Total Performance", 0.411, "7 LAST", "BB LLC", 21.531, "52x behind the leader; 6x behind mid-pack"],
    ["Financial Performance", 5.319, "7 LAST", "BB LLC", 37.687, "Operating loss + 25,000 shares. Next-worst (Bike Bros) is 12.856 — we are half of last-but-one"],
    ["Market Performance", 0.098, "6", "BB LLC", 0.345, "Only Bike Bros (0.085) is worse. 33% stock-outs; share x fill-rate"],
    ["Marketing Effectiveness", 0.738, "6", "Bike Bros / BB LLC", 0.765, "0.005 below average. Pack is 0.755–0.765; not a Q4 lever"],
    ["Investment in Future", 6.795, "2 STRENGTH", "Bike Bros", 7.401, "Do not chase. Leaders score LOW because revenue is in the denominator"],
    ["Wealth", 0.668, "7 LAST", "SpaceBikes", 1.017, "Cumulative losses. SpaceBikes is the only firm above 1.0"],
    ["Human Resource Management", 0.713, "5= (tie MILC)", "BB LLC", 0.767, "Only Bike Bros (0.691) is worse. Pay is low-middle, not last"],
    ["Asset Management", 0.353, "7 LAST", "BB LLC", 0.920, "$1.01M idle cash. Bike Bros 0.366 is almost as bad"],
    ["Manufacturing Productivity", 0.938, "2= (tie LiteCycle, Spoke'd Up)", "MILC Bikes", 0.995, "FALSE strength: we used 100% of the 8/day we scheduled. Does not punish stock-outs"],
    ["Financial Risk", 1.000, "1= (all firms)", "everyone", 1.000, "Nobody has debt. Zero differentiation"],
]
rank_start = r
r = block(ws, r, rows, wrap_col=6)
for i in range(1, 11):
    cell = ws.cell(rank_start + i, 3)
    read = str(cell.value or "")
    if "LAST" in read or read.startswith("6"):
        cell.fill = bad
    elif "STRENGTH" in read or read.startswith("2") or read.startswith("1"):
        cell.fill = good

# ---- two leagues ----
tp = SCORES["Total Performance"]
rows = [
    ["Two-league industry (Total Performance)", "Score", "x WeBike", "League"],
    ["BB LLC", 21.531, round(21.531 / 0.411, 1), "League 1 — profitable, high share, high asset turnover"],
    ["SpaceBikes", 10.753, round(10.753 / 0.411, 1), "League 1"],
    ["LiteCycle", 2.765, round(2.765 / 0.411, 1), "League 2 — mid-pack, tightly bunched 2.3–2.8"],
    ["MILC Bikes", 2.604, round(2.604 / 0.411, 1), "League 2"],
    ["Spoke'd Up", 2.316, round(2.316 / 0.411, 1), "League 2  (chart-read had ~1.0 — WRONG)"],
    ["Bike Bros", 0.900, round(0.900 / 0.411, 1), "League 3 — basement  (chart-read had ~2.5 — WRONG)"],
    ["WeBike", 0.411, 1.0, "League 3 — LAST"],
]
league_start = r
r = block(ws, r, rows, wrap_col=4)
for i in range(1, 8):
    if ws.cell(league_start + i, 1).value == "WeBike":
        for c in range(1, 5):
            ws.cell(league_start + i, c).fill = ours
            ws.cell(league_start + i, c).font = bold

# ---- product identity ----
prod = 1.0
for name in COMPONENTS:
    prod *= SCORES[name][WEBIKE]
rows = [
    ["Product identity (Total Performance = product of 9 indicators)", "Value"],
    ["Product of WeBike's 9 indicators", round(prod, 4)],
    ["Reported Total Performance", 0.411],
    ["Match?", "YES" if abs(prod - 0.411) < 0.002 else f"NO ({prod:.4f})"],
]
r = block(ws, r, rows)

# ---- one-at-a-time sensitivity ----
rows = [
    ["If WeBike matched the industry average on ONE indicator (others held)", "New TP", "x actual", "Note"],
]
actual_tp = 0.411
for name in COMPONENTS:
    ours_v = SCORES[name][WEBIKE]
    avg = sum(SCORES[name]) / 7
    new_tp = actual_tp * (avg / ours_v)
    note = ""
    if new_tp < actual_tp:
        note = "matching average would HURT us — this is a strength"
    rows.append([name, round(new_tp, 3), round(new_tp / actual_tp, 2), note])
rows.append(["All indicators at average (implied)", round(
    (lambda p: p)(
        (sum(SCORES["Total Performance"]) / 7)
    ), 3), round((sum(SCORES["Total Performance"]) / 7) / actual_tp, 1),
    "industry mean TP is 5.897 — 14x us. Not a realistic Q4 target"])
r = block(ws, r, rows, wrap_col=4)

# ---- path to mid-pack ----
rows = [
    ["Path to mid-pack (~2.5, LiteCycle/MILC/Spoke'd Up)", "TP", "x actual", "What it takes"],
    ["Q3 actual", 0.411, 1.00, "baseline"],
    ["Serve 100% of demand (Market 0.098 → 0.150)", round(0.411 / 0.098 * 0.150, 3),
     round((0.411 / 0.098 * 0.150) / 0.411, 2), "schedule OC to the forecast; no printers required"],
    ["+ Asset 0.353 → 0.55", round(0.411 / 0.098 * 0.150 / 0.353 * 0.55, 3),
     round((0.411 / 0.098 * 0.150 / 0.353 * 0.55) / 0.411, 2), "deploy idle cash into revenue (people, web, ads, city #3)"],
    ["+ Financial 5.319 → 14.8 (mid-pack)",
     round(0.411 / 0.098 * 0.150 / 0.353 * 0.55 / 5.319 * 14.8, 3),
     round((0.411 / 0.098 * 0.150 / 0.353 * 0.55 / 5.319 * 14.8) / 0.411, 2),
     "profitable quarter + stop issuing stock. This is the jump out of the basement"],
    ["Do NOT chase Investment in Future", None, None,
     "We are already 2nd. SpaceBikes is LAST on IIF (3.466) and 2nd overall. Bike Bros is BEST on IIF (7.401) and 6th overall"],
    ["Do NOT protect Manufacturing Productivity by under-scheduling", None, None,
     "0.938 is what you get when you use 100% of a too-small plan. Market Performance is the indicator that saw the 204 lost units"],
]
path_start = r
r = block(ws, r, rows, wrap_col=4)
ws.cell(path_start + 2, 2).fill = yellow  # serve 100%
ws.cell(path_start + 3, 2).fill = yellow  # + asset
ws.cell(path_start + 4, 2).fill = good    # + financial → mid-pack

# ---- chart-read corrections ----
rows = [
    ["Chart-read BSC in Q3_Industry_Graphs — CORRECTIONS", "Graph-read TP", "Actual TP", "Error"],
    ["Bike Bros", 2.5, 0.900, "was placed in mid-pack; actually in the basement with us"],
    ["Spoke'd Up", 1.0, 2.316, "was placed in the basement; actually mid-pack with LiteCycle/MILC"],
    ["LiteCycle", 2.0, 2.765, "directionally fine, a bit low"],
    ["MILC Bikes", 2.8, 2.604, "directionally fine"],
    ["SpaceBikes", 11.0, 10.753, "fine"],
    ["BB LLC", 22.0, 21.531, "fine"],
    ["WeBike", 0.411, 0.411, "exact — we had the live scorecard for ourselves"],
]
corr_start = r
r = block(ws, r, rows, wrap_col=4)
ws.cell(corr_start + 1, 1).fill = bad      # Bike Bros
ws.cell(corr_start + 2, 1).fill = yellow   # Spoke'd Up

for col, width in enumerate([52, 14, 14, 14, 14, 18, 16, 16, 14, 12, 12, 14], start=1):
    ws.column_dimensions[get_column_letter(col)].width = width
ws.column_dimensions["F"].width = 88
ws.column_dimensions["D"].width = 72

# ---- also stamp ranks onto existing Q3_BSC ----
if "Q3_BSC" in wb.sheetnames:
    bsc = wb["Q3_BSC"]
    bsc["H3"] = "WeBike rank"
    bsc["H3"].font = bold
    bsc["H3"].fill = hdr
    # rows 4-13 match the original indicator order
    order = [
        "Total Performance", "Financial Performance", "Market Performance",
        "Marketing Effectiveness", "Investment in Future", "Wealth",
        "Human Resource Mgmt", "Asset Management", "Manufacturing Productivity",
        "Financial Risk",
    ]
    # Q3_BSC uses a slightly different HR label
    label_map = {"Human Resource Mgmt": "Human Resource Management"}
    for i, label in enumerate(order):
        key = label_map.get(label, label)
        rk = webike_ranks[key]
        cell = bsc.cell(4 + i, 8, rk)
        if rk >= 6:
            cell.fill = bad
        elif rk <= 2:
            cell.fill = good

wb.save(DST)
print(f"Wrote Q3_Advanced_BSC to {DST}")
print(f"WeBike product of 9 indicators = {prod:.4f} (reported 0.411)")
print("WeBike ranks:")
for name, rk in webike_ranks.items():
    print(f"  {rk:>2}  {name:30s}  {SCORES[name][WEBIKE]}")
