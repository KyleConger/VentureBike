"""Add Q3 brand judgment intelligence (all 20 industry brands) to Q3Data.xlsx."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DST = "quarters/Q3/Q3Data.xlsx"
wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
bad = PatternFill("solid", fgColor="FFC7CE")
good = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFEB9C")
ours = PatternFill("solid", fgColor="FFF2CC")

if "Q3_Brand_Judgment" in wb.sheetnames:
    del wb["Q3_Brand_Judgment"]
ws = wb.create_sheet("Q3_Brand_Judgment")

ws["A1"] = "Brand Judgment — World Market, Q3 actual (all 20 industry brands)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("Score = segment appeal of the brand DESIGN (1 = no appeal). Exact figures from the sim report. "
            "A score of 1 means the brand is irrelevant to that segment.")
ws["A2"].alignment = Alignment(wrap_text=True)

# brand, company, rec, mtn, speed, targeted segment
brands = [
    ("MACH I.I", "Bike Bros", 12, 1, 77, "Speed"),
    ("TERRAMAX", "Bike Bros", 56, 73, 1, "Mountain"),
    ("Mach 0.6", "Bike Bros", 9, 1, 61, "Speed"),
    ("TERRAMean", "Bike Bros", 59, 72, 1, "Mountain"),
    ("MountainCruise1", "Bike Bros", 76, 1, 1, "Recreation"),
    ("LiteSpeed Pro+", "LiteCycle", 12, 1, 77, "Speed"),
    ("LiteSpeed+", "LiteCycle", 7, 1, 74, "Speed"),
    ("LiteTrail Pro", "LiteCycle", 58, 70, 1, "Mountain"),
    ("Hike Bike", "WeBike", 56, 73, 1, "Mountain"),
    ("Swift Bike", "WeBike", 10, 1, 72, "Speed"),
    ("Blu Ruged Ballz", "BB LLC", 56, 73, 1, "Mountain"),
    ("Blu Aero Ballz", "BB LLC", 7, 1, 74, "Speed"),
    ("Blu Tail Ballz", "BB LLC", 58, 70, 1, "Mountain"),
    ("Blu Tube Ballz", "BB LLC", 9, 1, 76, "Speed"),
    ("Spoke'd Easy", "Spoke'd Up", 73, 1, 1, "Recreation"),
    ("Spoke'd Speed", "Spoke'd Up", 10, 1, 75, "Speed"),
    ("The Armstrong", "SpaceBikes", 9, 1, 76, "Speed"),
    ("Mars Rover", "SpaceBikes", 73, 1, 1, "Recreation"),
    ("Whole MILC MKII", "MILC Bikes", 74, 1, 1, "Recreation"),
    ("Skim MILC MKII", "MILC Bikes", 9, 1, 76, "Speed"),
]

r = 4
head = ["Brand", "Company", "Recreation", "Mountain", "Speed", "Targets", "Best score", "Gap to segment best"]
for c, v in enumerate(head, start=1):
    cell = ws.cell(r, c, v)
    cell.fill = hdr
    cell.font = bold

MTN_BEST, SPD_BEST, REC_BEST = 73, 77, 76
r += 1
for name, co, rec, mtn, spd, tgt in brands:
    row = [name, co, rec, mtn, spd, tgt]
    best = {"Mountain": mtn, "Speed": spd, "Recreation": rec}[tgt]
    ceiling = {"Mountain": MTN_BEST, "Speed": SPD_BEST, "Recreation": REC_BEST}[tgt]
    row += [best, ceiling - best]
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r, c, v)
        if co == "WeBike":
            cell.fill = ours
            cell.font = bold
    r += 1

# ---- segment ceilings ----
r += 1
rows = [
    ["SEGMENT CEILINGS — what the best design achieves", "Best score", "Who holds it", "WeBike", "Gap"],
    ["Mountain", 73, "Hike Bike (WeBike), TERRAMAX (Bike Bros), Blu Ruged Ballz (BB LLC)", 73, 0],
    ["Speed", 77, "MACH I.I (Bike Bros), LiteSpeed Pro+ (LiteCycle)", 72, 5],
    ["Recreation", 76, "MountainCruise1 (Bike Bros)", "n/a (56 spillover)", "n/a"],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.fill = hdr
            cell.font = bold
ws.cell(r + 1, 4).fill = good
ws.cell(r + 1, 5).fill = good
ws.cell(r + 2, 4).fill = bad
ws.cell(r + 2, 5).fill = bad
r += len(rows) + 2

# ---- headline finding ----
rows = [
    ["HEADLINE — Hike Bike is tied for the BEST Mountain brand in the industry", None],
    ["Hike Bike Mountain judgment", 73],
    ["BB LLC Blu Ruged Ballz Mountain judgment", 73],
    ["Our Mountain share [%]", 22.1],
    ["BB LLC Mountain share [%]", 46.3],
    ["Share ratio on an IDENTICAL brand score", "=B{0}/B{1}".format(r + 4, r + 3)],
    ["Conclusion", "BB LLC takes 2.1x our Mountain share with a brand design scored EXACTLY the same as ours. "
     "The gap is therefore NOT product. It is supply (we stocked out 33%), sales coverage "
     "(21 sales people vs our 14, ~88 demand/store rep vs our 46) and ad volume (24 vs 19). "
     "Do not redesign Hike Bike. Feed it."],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.font = bold
            cell.fill = good
ws.cell(r + 6, 2).alignment = Alignment(wrap_text=True)
r += len(rows) + 2

# ---- Speed gap ----
rows = [
    ["SPEED GAP — Swift Bike is the weakest real Speed brand", "Rec", "Mtn", "Speed", "Note"],
    ["MACH I.I (Bike Bros)", 12, 1, 77, "co-best Speed design"],
    ["LiteSpeed Pro+ (LiteCycle)", 12, 1, 77, "co-best Speed design"],
    ["Blu Tube Ballz (BB LLC)", 9, 1, 76, None],
    ["The Armstrong (SpaceBikes)", 9, 1, 76, None],
    ["Skim MILC MKII (MILC)", 9, 1, 76, None],
    ["Spoke'd Speed (Spoke'd Up)", 10, 1, 75, None],
    ["LiteSpeed+ (LiteCycle)", 7, 1, 74, None],
    ["Blu Aero Ballz (BB LLC)", 7, 1, 74, None],
    ["Swift Bike (WeBike)", 10, 1, 72, "2nd WORST — 5 points off the ceiling"],
    ["Mach 0.6 (Bike Bros)", 9, 1, 61, "only brand worse than ours"],
    ["Read", None, None, None,
     "8 of 9 rival Speed brands beat Swift Bike. The winning profile (12/1/77) is nearly identical to ours "
     "(10/1/72) - so a modest redesign should close most of the 5-point gap. Combined with our Speed AD "
     "judgment of 70, weak product + weak ad explains 7.5% share in a 2,878-unit segment."],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.fill = hdr
            cell.font = bold
ws.cell(r + 9, 1).fill = bad
ws.cell(r + 9, 4).fill = bad
ws.cell(r + len(rows) - 1, 5).alignment = Alignment(wrap_text=True)
r += len(rows) + 2

# ---- design convergence ----
rows = [
    ["DESIGN CONVERGENCE — the market has found the optimal profiles", "Brands sharing it"],
    ["56 / 73 / 1  (Mountain optimum)", "TERRAMAX (Bike Bros), Hike Bike (WeBike), Blu Ruged Ballz (BB LLC)"],
    ["58 / 70 / 1  (Mountain, 2nd tier)", "LiteTrail Pro (LiteCycle), Blu Tail Ballz (BB LLC)"],
    ["12 / 1 / 77  (Speed optimum)", "MACH I.I (Bike Bros), LiteSpeed Pro+ (LiteCycle)"],
    ["9 / 1 / 76   (Speed, 2nd tier)", "Blu Tube Ballz (BB LLC), The Armstrong (SpaceBikes), Skim MILC MKII (MILC)"],
    ["7 / 1 / 74   (Speed, 3rd tier)", "LiteSpeed+ (LiteCycle), Blu Aero Ballz (BB LLC)"],
    ["73 / 1 / 1   (Recreation)", "Spoke'd Easy (Spoke'd Up), Mars Rover (SpaceBikes)"],
    ["Read", "Multiple firms have landed on byte-identical designs. Brand design is becoming COMMODITISED - "
     "no one can win on product alone. Nobody exceeds 73 in Mountain or 77 in Speed, so those look like "
     "hard ceilings. Differentiation must come from supply, channel, ads and price."],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.fill = hdr
            cell.font = bold
ws.cell(r + len(rows) - 1, 2).alignment = Alignment(wrap_text=True)
r += len(rows) + 2

# ---- portfolio breadth ----
rows = [
    ["PORTFOLIO BREADTH", "Brands", "Segments covered", "Overall share [%]", "Note"],
    ["Bike Bros", 5, "Rec + Mtn + Speed", 7.5, "MOST brands, LOWEST share - 1 store, 0 web centers"],
    ["BB LLC", 4, "Rec* + Mtn + Speed", 23.0, "leader; 2 Mtn + 2 Speed brands, Rec via spillover only"],
    ["LiteCycle", 3, "Mtn + Speed", 15.0, None],
    ["WeBike", 2, "Mtn + Speed", 9.6, "at parity on count"],
    ["Spoke'd Up", 2, "Rec + Speed", 13.0, None],
    ["SpaceBikes", 2, "Rec + Speed", 19.8, "cumulative LEADER on just 2 brands - skips Mountain"],
    ["MILC Bikes", 2, "Rec + Speed", 12.2, None],
    ["Read", None, None, None,
     "Brand COUNT does not drive share: Bike Bros has 5 brands and last place; SpaceBikes has 2 and leads "
     "cumulatively. A 3rd brand is NOT our priority - filling demand for the 2 we have is."],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.fill = hdr
            cell.font = bold
ws.cell(r + 4, 1).fill = ours
ws.cell(r + len(rows) - 1, 5).alignment = Alignment(wrap_text=True)
r += len(rows) + 2

# ---- Recreation spillover ----
rows = [
    ["RECREATION SPILLOVER — a free option we are not using", "Score"],
    ["Hike Bike appeal to Recreation buyers", 56],
    ["Best dedicated Recreation brand (MountainCruise1)", 76],
    ["Recreation segment size [units]", 2060],
    ["Recreation searches (highest of 3 segments)", 1350],
    ["Our Recreation share [%]", 2.6],
    ["Our Recreation demand [units, untargeted]", 54],
    ["Read", "Hike Bike already scores 56 with Recreation buyers and pulled 54 units we never asked for. "
     "That is real latent demand in the LARGEST-search segment. But 56 vs 76 means we cannot win Recreation "
     "on spillover - it would need a dedicated brand. Park this as a Q5 option, not a Q4 one."],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.fill = hdr
            cell.font = bold
ws.cell(r + len(rows) - 1, 2).alignment = Alignment(wrap_text=True)
r += len(rows) + 2

# ---- Q4 actions ----
rows = [
    ["Q4 BRAND ACTIONS", "Decision"],
    ["Hike Bike (Mountain)", "DO NOT REDESIGN. Tied for best in industry at 73 - already at the ceiling. "
     "Every dollar should go to supplying and selling it, not improving it."],
    ["Swift Bike (Speed)", "REDESIGN toward the 12/1/77 profile. We are 5 points off the ceiling and 8 of 9 "
     "rival Speed brands beat us. Also strengthen the Speed ad (judgment 70)."],
    ["Third brand", "NO. Bike Bros proves brand count does not buy share. Fix supply first."],
    ["Recreation", "DEFER. Spillover of 56 is not competitive vs 76; revisit in Q5 with a dedicated design."],
    ["Price", "HOLD. Our Mountain brand matches the best design and BB LLC charges ~$1,480 - we have "
     "room at $1,365, not a reason to discount."],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.fill = hdr
            cell.font = bold
        elif c == 2:
            cell.fill = yellow
            cell.alignment = Alignment(wrap_text=True)

ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 20
ws.column_dimensions["E"].width = 70
ws.column_dimensions["F"].width = 14
ws.column_dimensions["G"].width = 12
ws.column_dimensions["H"].width = 20

wb.save(DST)
print("Added Q3_Brand_Judgment")
