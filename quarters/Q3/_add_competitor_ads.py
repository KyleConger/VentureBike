"""Add Q3 competitor-ad COPY (ranked benefits) and correct the Swift Bike diagnosis.

Judgment scores already live on Q3_Ad_Judgment. This sheet is the missing
ingredient: what each ad actually claimed, in rank order.
"""
from collections import Counter
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DST = "quarters/Q3/Q3Data.xlsx"

# Short key -> Workspace label (exact)
LABEL = {
    "brand": "Mention brand name",
    "local": "Local sales & service",
    "best_rec": "Highest rated Recreation bike",
    "best_mtn": "Highest rated Mountain bike",
    "best_spd": "Highest rated Speed bike",
    "greenway": "Picture of riders on a greenway",
    "trail": "Picture of a rider on a steep trail",
    "race": "Picture of road race",
    "lights": "Ride safely after dark with lights",
    "reflectors": "Added safety - reflectors everywhere!",
    "carry": "Carry your stuff easily",
    "gears": "Tackle steep climbs with more gears",
    "price": "Carbon fiber quality at a great price",
    "print3d": "A tailor-made bike just for you! 3D printing",
    "elite": "Elite look - a ride of distinction",
    "light": "Enjoy your ride - carbon fiber light",
    "adventure": "Have an adventure on a carbon bike!",
    "tough": "Go anywhere on a tough carbon bike",
    "mountains": "Mountains are no longer difficult",
    "gym": "The gym is wherever you take it",
    "wind": "Ride a wind-cheater",
    "fast": "Fast and furious",
    "fun": "Lots of fun to ride",
    "tread": "Grab the path with high tread tires",
    "racing": "Roll fast with racing tires",
    "seat": "Easy riding on a comfort seat",
    "bars": "Keep your head high - comfort straight handlebar",
}

# ad, company, brand, segment, rec, mtn, speed, {benefit: rank}
ADS = [
    ("Excluspeedity", "Bike Bros", "MACH I.I", "Speed", 0, 12, 72,
     {"race": 1, "brand": 2, "wind": 3, "fast": 4, "racing": 5, "light": 6}),
    ("TerraTech", "Bike Bros", "TERRAMAX", "Mountain", 23, 78, 30,
     {"trail": 1, "brand": 2, "tough": 3, "local": 4, "mountains": 5, "gears": 6}),
    ("Easy Rider", "Bike Bros", "MountainCruise1", "Recreation", 75, 25, 17,
     {"seat": 1, "brand": 2, "bars": 3, "greenway": 4, "local": 5, "fun": 6}),
    ("Mountainability", "Bike Bros", "TERRAMAX", "Mountain", 35, 77, 32,
     {"trail": 1, "brand": 2, "adventure": 3, "gears": 4, "tread": 5, "local": 6}),
    ("AndStill", "Bike Bros", "MACH I.I", "Speed", 22, 24, 78,
     {"race": 1, "brand": 2, "wind": 3, "elite": 4, "print3d": 5, "racing": 6, "local": 7}),
    ("Speed of Lite 1", "LiteCycle", "LiteSpeed Pro+", "Speed", 10, 22, 77,
     {"race": 1, "brand": 2, "wind": 3, "elite": 4, "racing": 5, "print3d": 6, "fast": 7}),
    ("Speed of Lite 2", "LiteCycle", "LiteSpeed+", "Speed", 16, 13, 73,
     {"race": 1, "brand": 2, "light": 3, "wind": 4, "racing": 5, "elite": 6}),
    ("Trail Blazing 1", "LiteCycle", "LiteTrail Pro", "Mountain", 18, 79, 34,
     {"trail": 1, "brand": 2, "mountains": 3, "adventure": 4, "gears": 5, "tread": 6}),
    ("HikeBike 1", "WeBike", "Hike Bike", "Mountain", 33, 80, 26,
     {"brand": 1, "trail": 2, "gears": 3, "tread": 4, "adventure": 5, "tough": 6, "mountains": 7}),
    ("Swift Bike", "WeBike", "Swift Bike", "Speed", 52, 36, 70,
     {"race": 1, "light": 2, "price": 3, "brand": 4, "lights": 5, "reflectors": 6}),
    ("BB Big Momma", "BB LLC", "Blu Ruged Ballz", "Mountain", 28, 81, 27,
     {"gears": 1, "tread": 2, "best_mtn": 3, "local": 4, "brand": 5, "trail": 6}),
    ("Unleash lil pap", "BB LLC", "Blu Tube Ballz", "Speed", 15, 16, 76,
     {"wind": 1, "racing": 2, "elite": 3, "local": 4, "brand": 5, "race": 6}),
    ("Spoke'd Easy", "Spoke'd Up", "Spoke'd Easy", "Recreation", 59, 27, 21,
     {"brand": 1, "seat": 2, "fun": 3, "bars": 4, "greenway": 5, "light": 6,
      "price": 7, "carry": 8, "best_rec": 9}),
    ("Spoke'd Speed", "Spoke'd Up", "Spoke'd Speed", "Speed", 4, 18, 57,
     {"brand": 1, "fast": 2, "wind": 3, "racing": 4, "light": 5, "race": 6,
      "elite": 7, "gym": 8, "print3d": 9}),
    ("The Armstrong 1", "SpaceBikes", "The Armstrong", "Speed", 2, 21, 77,
     {"best_spd": 1, "race": 2, "fast": 3, "brand": 4, "wind": 5, "local": 6}),
    ("Mars Rover 1", "SpaceBikes", "Mars Rover", "Recreation", 79, 25, 24,
     {"best_rec": 1, "fun": 2, "greenway": 3, "brand": 4, "seat": 5, "reflectors": 6, "local": 7}),
    ("Skim MILC MKII", "MILC Bikes", "Skim MILC MKII", "Speed", 10, 19, 75,
     {"race": 1, "brand": 2, "wind": 3, "elite": 4, "racing": 5, "light": 6, "fast": 7}),
    ("Whole MILC MKII", "MILC Bikes", "Whole MILC MKII", "Recreation", 77, 47, 41,
     {"seat": 1, "brand": 2, "bars": 3, "light": 4, "price": 5, "local": 6}),
]


def ordered(ranks):
    return [k for k, _ in sorted(ranks.items(), key=lambda kv: kv[1])]


def rank_line(ranks):
    return " -> ".join(f"{i}. {LABEL[k]}" for i, k in enumerate(ordered(ranks), start=1))


print("RANK VALIDATION")
bad = []
for ad, co, br, seg, rec, mtn, spd, ranks in ADS:
    vals = sorted(ranks.values())
    expect = list(range(1, len(vals) + 1))
    ok = vals == expect
    if not ok:
        bad.append((ad, vals))
    print(f"  {'OK' if ok else 'GAP'}  {ad:<17} {len(ranks)} benefits  {rank_line(ranks)[:80]}")
assert not bad, bad
print("  all 18 ads have consecutive ranks 1..n\n")

spd_rivals = [a for a in ADS if a[3] == "Speed" and a[1] != "WeBike"]
mtn_rivals = [a for a in ADS if a[3] == "Mountain" and a[1] != "WeBike"]
swift = next(a for a in ADS if a[0] == "Swift Bike")
hike = next(a for a in ADS if a[0] == "HikeBike 1")

print("SPEED BENEFIT FREQUENCY (8 rival ads)")
freq = Counter()
for *_, ranks in spd_rivals:
    freq.update(ranks.keys())
swift_set = set(swift[7])
for k, n in freq.most_common():
    flag = "  SWIFT HAS" if k in swift_set else ""
    unique = "  <-- SWIFT ONLY among Speed" if n == 0 else ""
    print(f"  {n}/8  {LABEL[k]}{flag}")
print("  SWIFT-ONLY (0/8 rivals):")
for k in ordered(swift[7]):
    if freq[k] == 0:
        print(f"    {LABEL[k]}  rank {swift[7][k]}")

print("\nMOUNTAIN BENEFIT FREQUENCY (4 rival ads)")
mf = Counter()
for *_, ranks in mtn_rivals:
    mf.update(ranks.keys())
hike_set = set(hike[7])
for k, n in mf.most_common():
    flag = "  HIKE HAS" if k in hike_set else "  HIKE MISSING"
    print(f"  {n}/4  {LABEL[k]}{flag}")

wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
badf = PatternFill("solid", fgColor="FFC7CE")
good = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFEB9C")
ours = PatternFill("solid", fgColor="FFF2CC")

if "Q3_Competitor_Ads" in wb.sheetnames:
    del wb["Q3_Competitor_Ads"]
ws = wb.create_sheet("Q3_Competitor_Ads")

ws["A1"] = "Competitor Ads — ranked benefits, Q3 actual (all 18 ads)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("Source: Competitors' Ads research, previous quarter (Q3). Rank 1 = most visible. "
            "HikeBike 1 dropped the Q2 'great price' claim (now 7 benefits). "
            "Swift Bike is the only Speed ad that claims lights, reflectors, or great price.")
ws["A2"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A2:I2")
ws.row_dimensions[2].height = 36


def block(start, rows, wrap_col=None):
    for i, row in enumerate(rows):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(start + i, c, v)
            if i == 0:
                cell.fill = hdr
                cell.font = bold
            elif wrap_col and c == wrap_col:
                cell.alignment = Alignment(wrap_text=True)
    return start + len(rows) + 2


r = 4
rows = [["Ad", "Company", "Brand advertised", "Segment", "Rec", "Mtn", "Speed",
         "n benefits", "Ranked benefits (1 = most visible)"]]
for ad, co, br, seg, rec, mtn, spd, ranks in ADS:
    rows.append([ad, co, br, seg, rec, mtn, spd, len(ranks), rank_line(ranks)])
r2 = block(r, rows, wrap_col=9)
for i in range(1, len(rows)):
    ws.cell(r + i, 9).alignment = Alignment(wrap_text=True)
    if rows[i][1] == "WeBike":
        for c in range(1, 10):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
        if rows[i][0] == "Swift Bike":
            ws.cell(r + i, 7).fill = badf
        else:
            ws.cell(r + i, 6).fill = good
    ws.row_dimensions[r + i].height = 36
r = r2

rows = [["SWIFT BIKE vs THE TWO BEST SPEED ADS", "AndStill (78, best)", "Speed of Lite 1 (77)", "Swift Bike (70)"]]
andstill = next(a for a in ADS if a[0] == "AndStill")
sol1 = next(a for a in ADS if a[0] == "Speed of Lite 1")
keys_show = ["race", "brand", "wind", "elite", "racing", "print3d", "fast", "local",
             "light", "price", "lights", "reflectors"]
rows.append(["Score in Speed", 78, 77, 70])
rows.append(["Off-target leakage (Rec+Mtn)", 46, 32, 88])
for k in keys_show:
    rows.append([LABEL[k],
                 andstill[7].get(k, "—"),
                 sol1[7].get(k, "—"),
                 swift[7].get(k, "—")])
r2 = block(r, rows)
for i in range(1, len(rows)):
    ws.cell(r + i, 4).fill = ours
    if rows[i][0] in (LABEL["price"], LABEL["lights"], LABEL["reflectors"]):
        ws.cell(r + i, 4).fill = badf
    if rows[i][0] in (LABEL["wind"], LABEL["elite"], LABEL["racing"], LABEL["print3d"], LABEL["fast"], LABEL["local"]):
        if rows[i][4 - 1] == "—":
            ws.cell(r + i, 4).fill = badf
r = r2

rows = [
    ["SPEED CONSENSUS (8 rival ads) vs SWIFT", "Rivals using it", "Typical rank", "Swift Bike"],
    ["Picture of road race", "8/8", "1 (five ads lead with it)", "HAS — rank 1 (the one thing we got right)"],
    ["Mention brand name", "8/8", "2 (five ads)", "rank 4 — too late; SERP top-3 never names the brand"],
    ["Ride a wind-cheater", "8/8", "3", "MISSING"],
    ["Roll fast with racing tires", "7/8", "4–6", "MISSING (bike HAS racing tires)"],
    ["Elite look - a ride of distinction", "6/8", "3–4", "MISSING"],
    ["Fast and furious", "5/8", "2–7", "MISSING"],
    ["Enjoy your ride - carbon fiber light", "4/8", "3–6", "HAS — rank 2"],
    ["Local sales & service", "3/8", "4–7", "MISSING (we are 80.9% store — should claim it)"],
    ["A tailor-made bike just for you! 3D printing", "3/8", "5–6 on the two best ads", "MISSING (we do 3D print)"],
    ["Ride safely after dark with lights", "0/8", "none", "HAS — rank 5. Rec cue. Explains Rec leakage 52."],
    ["Added safety - reflectors everywhere!", "0/8", "none", "HAS — rank 6. Rec cue."],
    ["Carbon fiber quality at a great price", "0/8", "none", "HAS — rank 3. Speed is the least price-sensitive segment; we are about to RAISE price to $1,580."],
]
r2 = block(r, rows, wrap_col=4)
for i, key in enumerate(
        ["ok", "late", "miss", "miss", "miss", "miss", "ok", "miss", "miss", "wrong", "wrong", "wrong"], start=1):
    if key == "miss":
        ws.cell(r + i, 4).fill = yellow
    elif key == "wrong":
        ws.cell(r + i, 4).fill = badf
    elif key == "ok":
        ws.cell(r + i, 4).fill = good
r = r2

rows = [
    ["DIAGNOSIS — the cloned-Mountain hypothesis is FALSE", "Finding"],
    ["What we thought", "Swift Bike inherited Hike Bike's messaging (steep trail / gears / tread / adventure)."],
    ["What the copy actually is",
     "1. road race  2. carbon light  3. great price  4. brand  5. lights  6. reflectors. "
     "Zero Mountain benefits. It is a Recreation-safety + price ad with a Speed picture."],
    ["Why Rec leakage is 52 (highest of any Speed ad)",
     "Lights + reflectors are Rec safety claims. Great price is a Rec/value claim. "
     "True on the bike (Speed recipe includes both) but they pull the wrong segment."],
    ["Why the score is 70 not 57",
     "We at least lead with the road-race picture. Spoke'd Speed has every Speed cue "
     "but ranks road race 6th, pads to 9 benefits (incl. gym), and scores 57. Ranking and clutter matter."],
    ["SERP organic (first 3 benefits only)",
     "Today: road race / carbon light / GREAT PRICE. After a consensus rebuild: road race / brand / wind-cheater."],
    ["HikeBike 1 vs Q2 lock",
     "Q2 lock had 8 benefits, last = great price. Q3 actual dropped that claim (7 benefits). "
     "Ranks 1–7 unchanged. Score 80. Leave it."],
]
r2 = block(r, rows, wrap_col=2)
ws.cell(r + 3, 2).fill = yellow
r = r2

rows = [
    ["Q4 SWIFT BIKE AD — DRAFT, NOT LOCKED (Unleash SERP snippet + AndStill extras)", "Rank"],
    ["Mention brand name (Ad/Brand columns; not a snippet)", 1],
    ["Ride a wind-cheater", 2],
    ["Roll fast with racing tires", 3],
    ["Elite look - a ride of distinction", 4],
    ["A tailor-made bike just for you! 3D printing", 5],
    ["Local sales & service", 6],
    ["Picture of road race", 7],
    ["DROP", "lights, reflectors, great price, carbon-light-as-#2"],
    ["Why this order", "Brand does not eat a snippet slot, so ranks 2-4 = Unleash's #1 Speed SERP "
     "(wind / racing / elite, 200 clicks). Picture last keeps print coverage without burying the snippet. "
     "AndStill (picture-first) scores 78 judgment but 7th organically."],
    ["Do not add a 2nd Swift ad", "Quality > count. Preview the score in the designer before committing."],
    ["HikeBike 1", "KEEP. Optional 1-pt chase: add 'Highest rated Mountain bike' (true, tied at 73) "
     "and/or local sales & service — BB Big Momma uses both and scores 81. Not worth a redesign fee unless the preview moves."],
]
r2 = block(r, rows, wrap_col=2)
for i in range(1, 8):
    ws.cell(r + i, 2).fill = good
ws.cell(r + 8, 2).fill = badf
for i in (9, 10, 11):
    ws.cell(r + i, 2).fill = yellow
    ws.cell(r + i, 2).alignment = Alignment(wrap_text=True)
r = r2

ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 28
ws.column_dimensions["D"].width = 88
ws.column_dimensions["E"].width = 12
ws.column_dimensions["F"].width = 12
ws.column_dimensions["G"].width = 12
ws.column_dimensions["H"].width = 14
ws.column_dimensions["I"].width = 88

# ---- patch leftover clone-hypothesis wording if an old sheet is present ----
if "Q3_Ad_Judgment" in wb.sheetnames:
    aj = wb["Q3_Ad_Judgment"]
    for row in aj.iter_rows(min_row=1, max_row=aj.max_row, max_col=3):
        blob = " ".join(str(cell.value) for cell in row if cell.value)
        if "inherited Hike Bike" in blob or "cloned Hike Bike" in blob:
            row[2].value = (
                "FALSE: Swift Bike did NOT clone Hike Bike's Mountain messaging. Actual ranks: "
                "road race -> carbon light -> great price -> brand -> lights -> reflectors. "
                "Zero Mountain claims. Rec-safety + price ad with a Speed picture. "
                "See Q3_Competitor_Ads."
            )
            row[2].fill = yellow
            row[2].alignment = Alignment(wrap_text=True)
        if "clone AndStill" in blob and "Unleash" not in blob:
            target = row[1] if row[1].value and "clone AndStill" in str(row[1].value) else row[2]
            target.value = (
                "REDESIGN: Unleash snippet + AndStill extras. 1 brand (not a snippet)  "
                "2 wind-cheater  3 racing tires  4 elite look  5 3D printing  "
                "6 local sales  7 road-race picture. DROP lights, reflectors, great price. "
                "Not locked. Detail: Q3_Competitor_Ads."
            )

wb.save(DST)
print("\nAdded Q3_Competitor_Ads and patched Q3_Ad_Judgment diagnosis")
