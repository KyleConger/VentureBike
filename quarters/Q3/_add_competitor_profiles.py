"""Add exact Q3 Competitors' Profiles report to Q3Data.xlsx.

Source: Workspace > Competitors' Profiles, results of Quarter 3.
These numbers SUPERSEDE chart-read approximations in Q3_Industry_Graphs.
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DST = "quarters/Q3/Q3Data.xlsx"

# Column order matches the Workspace report
FIRMS = [
    "Bike Bros",
    "LiteCycle",
    "WeBike",
    "BB LLC",
    "Spoke'd Up",
    "SpaceBikes",
    "MILC Bikes",
]

# metric -> 7 values in FIRMS order
RAW = {
    "BSC (Q3 quarterly, rounded)": [1, 3, 0, 22, 2, 11, 3],
    "Total demand": [490, 984, 627, 1507, 854, 1298, 799],
    "Number of brands": [5, 3, 2, 4, 2, 2, 2],
    "Average price": [1475, 1380, 1408, 1473, 1249, 1340, 1250],
    "Best Rec brand (>=70)": [
        "76 MountainCruise1", "N/A", "N/A", "N/A",
        "73 Spoke'd Easy", "73 Mars Rover", "74 Whole MILC MKII",
    ],
    "Best Mountain brand (>=70)": [
        "73 TERRAMAX", "70 LiteTrail Pro", "73 Hike Bike", "73 Blu Ruged Ballz",
        "N/A", "N/A", "N/A",
    ],
    "Best Speed brand (>=70)": [
        "77 MACH I.I", "77 LiteSpeed Pro+", "72 Swift Bike", "76 Blu Tube Ballz",
        "75 Spoke'd Speed", "76 The Armstrong", "76 Skim MILC MKII",
    ],
    "Regional inserts": [10, 12, 19, 24, 12, 17, 12],
    "Web pages": [5, 3, 2, 2, 2, 2, 2],
    "Organic SEM clicks": [213, 153, 144, 437, 82, 401, 187],
    "Organic SEM expenses": [5000, 3000, 2000, 2000, 2000, 2000, 2000],
    "Best Rec ad (>=70)": [
        "75 Easy Rider", "N/A", "N/A", "N/A",
        "N/A", "79 Mars Rover 1", "77 Whole MILC MKII",
    ],
    "Best Mountain ad (>=70)": [
        "78 TerraTech", "79 Trail Blazing 1", "80 HikeBike 1", "81 BB Big Momma",
        "N/A", "N/A", "N/A",
    ],
    "Best Speed ad (>=70)": [
        "78 AndStill", "77 Speed of Lite 1", "70 Swift Bike", "76 Unleash lil pap",
        "N/A", "77 The Armstrong 1", "75 Skim MILC MKII",
    ],
    "Stores": [1, 2, 2, 2, 2, 2, 1],
    "Store sales force": [7, 10, 11, 14, 11, 14, 7],
    "Web sales centers": [0, 1, 1, 1, 1, 1, 1],  # N/A -> 0
    "Web sales force": [0, 7, 3, 7, 3, 7, 7],
    "Web productivity features": [0, 4, 1, 4, 2, 4, 4],
    "Sales compensation": [22060, 25736, 24425, 27402, 24511, 24711, 27630],
    "Sales productivity [%]": [68, 73, 70, 78, 71, 74, 70],
    "Fixed capacity / day": [16, 16, 24, 24, 16, 32, 24],
    "Operating capacity / day": [15, 14, 8, 24, 15, 30, 17],
    "Overtime capacity / day": [0.00, 1.75, 1.00, 6.90, 1.88, 0.00, 0.00],
    "Worker compensation": [19609, 21622, 20757, 23042, 20757, 21004, 23293],
    "Worker productivity [%]": [70, 74, 72, 75, 72, 73, 73],
}

wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
bad = PatternFill("solid", fgColor="FFC7CE")
good = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFEB9C")
ours = PatternFill("solid", fgColor="FFF2CC")
thin = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

for name in ("Q3_Competitor_Profiles",):
    if name in wb.sheetnames:
        del wb[name]


def block(ws, start, rows, wrap_col=None, we_col=None):
    for i, row in enumerate(rows):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(start + i, c, v)
            cell.border = thin
            if i == 0:
                cell.fill = hdr
                cell.font = bold
            elif wrap_col and c == wrap_col:
                cell.alignment = Alignment(wrap_text=True)
            if we_col and i > 0 and c == we_col:
                cell.fill = ours
                cell.font = bold
    return start + len(rows) + 2


ws = wb.create_sheet("Q3_Competitor_Profiles")
ws["A1"] = "Competitors' Profiles — Q3 actual (EXACT from Workspace report)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = (
    "Source: Competitors' Profiles, results of previous quarter (Quarter 3). "
    "These numbers SUPERSEDE chart-read approximations in Q3_Industry_Graphs. "
    "Biggest correction: BB LLC fixed capacity is 24, not 16; they schedule all 24 "
    "plus 6.90 overtime. WeBike / BB LLC / MILC are tied at 24 printers; SpaceBikes owns 32."
)
ws["A2"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A2:H2")
ws.row_dimensions[2].height = 48

# ---- raw matrix (metrics x firms) ----
r = 4
rows = [["Metric"] + FIRMS]
for metric, vals in RAW.items():
    rows.append([metric] + list(vals))
we_col = FIRMS.index("WeBike") + 2
r = block(ws, r, rows, we_col=we_col)

# number formats
metric_names = list(RAW.keys())
for i, metric in enumerate(metric_names):
    row_i = 5 + i  # header at 4
    for c in range(2, 9):
        cell = ws.cell(row_i, c)
        if metric in ("Average price", "Sales compensation", "Worker compensation",
                      "Organic SEM expenses"):
            cell.number_format = "#,##0"
        elif metric in ("Total demand", "Organic SEM clicks"):
            cell.number_format = "#,##0"
        elif metric == "Overtime capacity / day":
            cell.number_format = "0.00"

# highlight our worst cells
our_row = {m: 5 + i for i, m in enumerate(metric_names)}
ws.cell(our_row["BSC (Q3 quarterly, rounded)"], we_col).fill = bad
ws.cell(our_row["Operating capacity / day"], we_col).fill = bad
ws.cell(our_row["Web productivity features"], we_col).fill = bad
ws.cell(our_row["Organic SEM clicks"], we_col).fill = bad
ws.cell(our_row["Best Speed brand (>=70)"], we_col).fill = yellow
ws.cell(our_row["Best Speed ad (>=70)"], we_col).fill = yellow
ws.cell(our_row["Best Mountain brand (>=70)"], we_col).fill = good
ws.cell(our_row["Best Mountain ad (>=70)"], we_col).fill = good

# ---- derived capacity ----
oc = RAW["Operating capacity / day"]
ot = RAW["Overtime capacity / day"]
fixed = RAW["Fixed capacity / day"]
demand = RAW["Total demand"]
wprod = RAW["Worker productivity [%]"]
sprod = RAW["Sales productivity [%]"]
store_sf = RAW["Store sales force"]
web_sf = RAW["Web sales force"]

cap_header = [
    "CAPACITY — derived (exact)", "Fixed / day", "Scheduled OC / day", "OT / day",
    "Effective (OC+OT)", "OC / fixed", "Demand",
]
cap_body = []
for i, co in enumerate(FIRMS):
    eff = oc[i] + ot[i]
    util = oc[i] / fixed[i]
    cap_body.append([co, fixed[i], oc[i], ot[i], round(eff, 2), util, demand[i]])
cap_body.sort(key=lambda x: -x[4])
rows = [cap_header] + cap_body
r2 = block(ws, r, rows)
for i in range(1, len(rows)):
    ws.cell(r + i, 4).number_format = "0.00"
    ws.cell(r + i, 5).number_format = "0.00"
    ws.cell(r + i, 6).number_format = "0.0%"
    ws.cell(r + i, 7).number_format = "#,##0"
    if rows[i][0] == "WeBike":
        for c in range(1, 8):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
        ws.cell(r + i, 3).fill = bad
        ws.cell(r + i, 5).fill = bad
        ws.cell(r + i, 6).fill = bad
    if rows[i][0] == "BB LLC":
        ws.cell(r + i, 2).fill = yellow
        ws.cell(r + i, 5).fill = good
r = r2

rows = [
    ["CHART-READ CORRECTIONS (Strategic Graphs were off)", "Chart-read", "Exact", "What changes"],
    ["BB LLC fixed capacity", 16, 24, "They bought the SAME 3 printers we did. 'Buys flexibility not printers' is FALSE."],
    ["BB LLC scheduled OC", 15, 24, "They schedule 100% of owned capacity, then add 6.90 OT on top."],
    ["MILC scheduled OC", 24, 17, "MILC owns 24 but schedules only 17 (70.8% of owned). Still well above us."],
    ["LiteCycle scheduled OC", 15, 14, "Off by 1; rank unchanged."],
    ["Spoke'd Up scheduled OC", 14, 15, "Off by 1; rank unchanged."],
    ["WeBike OC / OT / fixed", "8 / 1 / 24", "8 / 1.00 / 24", "Confirmed. Last on OC. Only firm that reduced OC in Q3 (from graphs)."],
    ["Fixed-capacity ranking", "SpaceBikes 32, WeBike/MILC 24, rest 16",
     "SpaceBikes 32; WeBike = BB LLC = MILC 24; rest 16",
     "Three firms own 24, not two. We are tied for 2nd, and we used the least."],
]
r2 = block(ws, r, rows, wrap_col=4)
for i in range(1, 4):
    ws.cell(r + i, 3).fill = yellow
ws.cell(r + 1, 4).fill = bad
r = r2

# BB LLC gold standard
bb_i = FIRMS.index("BB LLC")
bb_eff = oc[bb_i] + ot[bb_i]
bb_implied = bb_eff * (wprod[bb_i] / 100) * 65
we_i = FIRMS.index("WeBike")
we_eff = oc[we_i] + ot[we_i]
we_implied = we_eff * (wprod[we_i] / 100) * 65

rows = [
    ["BB LLC IS THE CAPACITY GOLD STANDARD", "Value", "Note"],
    ["BB LLC effective capacity (OC+OT)", round(bb_eff, 2), "Highest in the industry, slightly above SpaceBikes' 30"],
    ["Implied output at 75% x 65 days", round(bb_implied, 1), "Matches their 1,507 demand almost exactly — they scheduled TO demand"],
    ["WeBike implied output at 72% x 65 days", round(we_implied, 1), "Matches the 423 units we actually produced; 204 demand went unfilled"],
    ["SpaceBikes model", "32 owned, 30 scheduled, 0 OT", "Own more printers, no overtime. Served 1,298 with slack."],
    ["BB LLC model", "24 owned, 24 scheduled, 6.90 OT", "Same printers as us, sweat them with OT + 75% worker productivity."],
    ["WeBike model", "24 owned, 8 scheduled, 1.00 OT", "Same printers as the leader, used one-third of them."],
    ["Q4 read", None, "Matching BB LLC's 1,507 units on 24 printers is possible (full OC + ~7 OT + 75% productivity). "
     "Rio+web (scenario D, 1,171) fits inside 24 with a little OT. New York (scenario E, 1,571) still needs a 4th printer."],
]
r2 = block(ws, r, rows, wrap_col=3)
ws.cell(r + 1, 2).fill = good
ws.cell(r + 3, 2).fill = bad
ws.cell(r + 8, 3).fill = yellow
ws.cell(r + 8, 3).alignment = Alignment(wrap_text=True)
r = r2

# ---- pay vs productivity ----
sales_pay = RAW["Sales compensation"]
worker_pay = RAW["Worker compensation"]
sales_avg = sum(sales_pay) / 7
worker_avg = sum(worker_pay) / 7

rows = [[
    "COMPENSATION vs PRODUCTIVITY (exact)", "Sales pay", "Sales prod %",
    "vs sales avg", "Worker pay", "Worker prod %", "vs worker avg",
]]
for i, co in enumerate(FIRMS):
    rows.append([
        co, sales_pay[i], sprod[i], sales_pay[i] - sales_avg,
        worker_pay[i], wprod[i], worker_pay[i] - worker_avg,
    ])
rows_body = rows[1:]
rows_body.sort(key=lambda x: -x[2])
rows = [rows[0]] + rows_body
r2 = block(ws, r, rows)
for i in range(1, len(rows)):
    ws.cell(r + i, 2).number_format = "$#,##0"
    ws.cell(r + i, 4).number_format = '"+$"#,##0;"-$"#,##0'
    ws.cell(r + i, 5).number_format = "$#,##0"
    ws.cell(r + i, 7).number_format = '"+$"#,##0;"-$"#,##0'
    if rows[i][0] == "WeBike":
        for c in range(1, 8):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
        ws.cell(r + i, 4).fill = bad
        ws.cell(r + i, 7).fill = bad
r = r2

rows = [
    ["PAY FINDINGS", "Value", "Note"],
    ["Industry sales-pay average", round(sales_avg), "WeBike $24,425 is BELOW average — team goal 'pay > average' is currently failing"],
    ["Industry worker-pay average", round(worker_avg), "WeBike $20,757 is BELOW average (tied with Spoke'd Up; only Bike Bros lower)"],
    ["Sales productivity range", "68 – 78", "Ceiling is 78 (BB LLC), not ~75. Still nowhere near the 85% we planned on."],
    ["Worker productivity range", "70 – 75", "Ceiling is 75 (BB LLC). Industry is tight; 85% remains unreachable."],
    ["Highest sales pay", "MILC $27,630 at 70% prod", "Highest pay does NOT equal highest productivity. BB LLC $27,402 → 78%."],
    ["Highest worker pay", "MILC $23,293 at 73% prod", "BB LLC $23,042 → 75% (best productivity). LiteCycle $21,622 → 74%."],
    ["Q4 target to meet team goal", "Sales > $25,211 · Worker > $21,441",
     "Upper quartile is sales ~$27.4k (BB LLC/MILC) and worker ~$23.0k. Plan capacity on 73–75% workers, mid-70s sales — never 85%."],
]
r2 = block(ws, r, rows, wrap_col=3)
ws.cell(r + 1, 2).fill = bad
ws.cell(r + 2, 2).fill = bad
ws.cell(r + 3, 2).fill = yellow
r = r2

# ---- channel ----
rows = [[
    "CHANNEL (exact)", "Stores", "Store SF", "Web centres", "Web SF",
    "Web features / 4", "Total SF", "Demand", "Demand / SF",
]]
for i, co in enumerate(FIRMS):
    total_sf = store_sf[i] + web_sf[i]
    per = demand[i] / total_sf if total_sf else 0
    rows.append([
        co, RAW["Stores"][i], store_sf[i], RAW["Web sales centers"][i], web_sf[i],
        RAW["Web productivity features"][i], total_sf, demand[i], round(per, 1),
    ])
rows_body = rows[1:]
rows_body.sort(key=lambda x: -x[7])
rows = [rows[0]] + rows_body
r2 = block(ws, r, rows)
for i in range(1, len(rows)):
    ws.cell(r + i, 8).number_format = "#,##0"
    ws.cell(r + i, 9).number_format = "0.0"
    if rows[i][0] == "WeBike":
        for c in range(1, 10):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
        ws.cell(r + i, 6).fill = bad
        ws.cell(r + i, 9).fill = bad
r = r2

rows = [
    ["CHANNEL FINDINGS", "Value", "Note"],
    ["Store counts", "5 firms have 2; Bike Bros AND MILC have 1",
     "City #3 is still unclaimed. MILC is web-first (confirmed 1 store + 7 web staff + 4 features)."],
    ["Web feature pattern", "4 tactics → 7 web staff; 1–2 tactics → 3 web staff",
     "Confirmed. WeBike 1 feature / 3 staff; Spoke'd Up 2 / 3; leaders 4 / 7."],
    ["Demand per sales person", "WeBike 44.8 = LAST",
     "BB LLC 71.8 · Bike Bros 70.0 (one store!) · SpaceBikes 61.8 · Spoke'd Up 61.0"],
    ["Headcount is NOT at parity", "Leaders have 21 (BB LLC, SpaceBikes)",
     "WeBike / Spoke'd Up / MILC = 14. LiteCycle 17. The 'parity at 14' line was wrong — leaders run 50% more people."],
]
r2 = block(ws, r, rows, wrap_col=3)
ws.cell(r + 3, 2).fill = bad
ws.cell(r + 4, 3).fill = yellow
r = r2

# ---- marketing ----
sem = RAW["Organic SEM clicks"]
sem_exp = RAW["Organic SEM expenses"]
inserts = RAW["Regional inserts"]
rows = [[
    "MARKETING (exact)", "Inserts", "SEM clicks", "SEM $", "Clicks / $1k",
    "Web pages", "Best Mtn ad", "Best Speed ad",
]]
for i, co in enumerate(FIRMS):
    cpk = sem[i] / (sem_exp[i] / 1000) if sem_exp[i] else 0
    rows.append([
        co, inserts[i], sem[i], sem_exp[i], round(cpk, 1),
        RAW["Web pages"][i],
        RAW["Best Mountain ad (>=70)"][i],
        RAW["Best Speed ad (>=70)"][i],
    ])
rows_body = rows[1:]
rows_body.sort(key=lambda x: -x[2])
rows = [rows[0]] + rows_body
r2 = block(ws, r, rows)
for i in range(1, len(rows)):
    ws.cell(r + i, 3).number_format = "#,##0"
    ws.cell(r + i, 4).number_format = "$#,##0"
    ws.cell(r + i, 5).number_format = "0.0"
    if rows[i][0] == "WeBike":
        for c in range(1, 9):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
r = r2

rows = [
    ["MARKETING FINDINGS", "Value", "Note"],
    ["SEM is a quality game, not a budget game",
     "5 of 7 firms spend $2,000",
     "BB LLC 437 clicks and SpaceBikes 401 clicks on the SAME $2,000 we spent for 144. "
     "Bike Bros spent $5,000 for 213 clicks — worse ROI. Pages don't explain it (BB LLC has 2 pages, we have 2)."],
    ["Mountain ad", "HikeBike 1 = 80, 2nd of 4", "BB Big Momma 81. One point off the top. Do not redesign."],
    ["Speed ad", "Swift Bike = 70, last of scored ads",
     "Spoke'd Up Speed ad is below 70 (N/A). We are last among anyone who cleared the bar. Redesign."],
    ["Average price", "WeBike $1,408 = 3rd",
     "Bike Bros $1,475 edges BB LLC $1,473. 'BB LLC highest price' was a $2 rounding error. Rec specialists (Spoke'd Up $1,249, MILC $1,250) sit ~$160 below us."],
    ["Brand count ≠ share", "Bike Bros 5 brands, 490 demand, BSC 1",
     "SpaceBikes 2 brands, 1,298 demand, BSC 11. BB LLC 4 brands (2 Mtn + 2 Speed, no Rec >=70). A 3rd brand is still not a Q4 priority."],
    ["BB LLC Rec = N/A", "No Rec brand or ad >= 70",
     "'All 3 segments' is spillover (99 Rec units), not a Rec franchise. They win Mountain + Speed only."],
]
r2 = block(ws, r, rows, wrap_col=3)
ws.cell(r + 1, 3).fill = yellow
ws.cell(r + 3, 2).fill = bad
r = r2

# Q3 quarterly BSC
rows = [
    ["Q3 QUARTERLY BSC (this report, integer)", "Score", "Demand", "Read"],
    ["BB LLC", 22, 1507, "Q3 single-quarter leader by a mile"],
    ["SpaceBikes", 11, 1298, "Strong Q3 too — cumulative lead plus a double-digit quarter"],
    ["LiteCycle", 3, 984, "Tied third"],
    ["MILC Bikes", 3, 799, "Tied third"],
    ["Spoke'd Up", 2, 854, None],
    ["Bike Bros", 1, 490, "5 brands, 1 store, 0 web — last among firms that sold anything much"],
    ["WeBike", 0, 627, "0.411 Total Performance rounds to 0 on this report. Last."],
]
r2 = block(ws, r, rows, wrap_col=4)
ws.cell(r + 1, 2).fill = good
ws.cell(r + 7, 2).fill = bad
for i in range(1, 8):
    ws.cell(r + i, 3).number_format = "#,##0"
    if rows[i][0] == "WeBike":
        for c in range(1, 5):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
        ws.cell(r + i, 2).fill = bad
r = r2

ws.column_dimensions["A"].width = 58
for col in "BCDEFGH":
    ws.column_dimensions[col].width = 22
ws.column_dimensions["D"].width = 28

# ---- patch Q3_Industry_Graphs with a correction banner if the sheet exists ----
if "Q3_Industry_Graphs" in wb.sheetnames:
    g = wb["Q3_Industry_Graphs"]
    # find a free row near the top note
    g["A3"] = (
        "SUPERSEDED IN PART by Q3_Competitor_Profiles (exact report). "
        "Corrections: BB LLC fixed=24 OC=24 OT=6.90 · MILC OC=17 not 24 · "
        "LiteCycle OC=14 · Spoke'd Up OC=15 · sales prod ceiling 78 (BB LLC) · "
        "sales pay avg $25,211 / worker avg $21,441 — we are BELOW both."
    )
    g["A3"].font = Font(bold=True, color="9C0006")
    g["A3"].alignment = Alignment(wrap_text=True)
    g.merge_cells("A3:F3")
    g.row_dimensions[3].height = 36
    g["A3"].fill = yellow

wb.save(DST)

# stdout summary for the chat
print("Added Q3_Competitor_Profiles")
print()
print("CAPACITY (exact, ranked by OC+OT)")
for i in sorted(range(7), key=lambda i: -(oc[i] + ot[i])):
    print(f"  {FIRMS[i]:<12} fixed {fixed[i]:>2}  OC {oc[i]:>2}  OT {ot[i]:4.2f}  "
          f"eff {oc[i]+ot[i]:5.2f}  util {oc[i]/fixed[i]:5.1%}  demand {demand[i]:>5}")
print()
print(f"Sales pay avg ${sales_avg:,.0f}  WeBike ${sales_pay[we_i]:,}  gap ${sales_pay[we_i]-sales_avg:,.0f}")
print(f"Worker pay avg ${worker_avg:,.0f}  WeBike ${worker_pay[we_i]:,}  gap ${worker_pay[we_i]-worker_avg:,.0f}")
print(f"Sales prod range {min(sprod)}-{max(sprod)}  Worker prod range {min(wprod)}-{max(wprod)}")
print(f"BB LLC implied units {bb_implied:.1f} vs demand {demand[bb_i]}")
print(f"WeBike implied units {we_implied:.1f} vs produced ~423 / demand {demand[we_i]}")
