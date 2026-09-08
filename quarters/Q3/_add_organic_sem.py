"""Add Q3 organic SEM listings, verify CTR vs search volume, and size the web-ops gap."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DST = "quarters/Q3/Q3Data.xlsx"

SEARCHES = {"Recreation": 1333, "Mountain": 1027, "Speed": 894}

# segment, pos, company, ad, brand, line1, line2, line3, clicks, ctr, ad_j, page_inv
SEM = [
    ("Recreation", 1, "SpaceBikes", "Mars Rover 1", "Mars Rover",
     "Highest rated Recreation bike", "Lots of fun to ride",
     "Picture of riders on a greenway", 303, 0.2273, 79, 1000),
    ("Recreation", 2, "MILC Bikes", "Whole MILC MKII", "Whole MILC MKII",
     "Easy riding on a comfort seat", "Keep your head high - comfort straight handlebar",
     "Enjoy your ride - carbon fiber light", 144, 0.1080, 77, 1000),
    ("Recreation", 3, "Bike Bros", "Easy Rider", "MountainCruise1",
     "Easy riding on a comfort seat", "Keep your head high - comfort straight handlebar",
     "Picture of riders on a greenway", 88, 0.0660, 75, 1000),
    ("Recreation", 4, "Spoke'd Up", "Spoke'd Easy", "Spoke'd Easy",
     "Easy riding on a comfort seat", "Lots of fun to ride",
     "Keep your head high - comfort straight handlebar", 61, 0.0458, 59, 1000),
    ("Mountain", 1, "BB LLC", "BB Big Momma", "Blu Ruged Ballz",
     "Tackle steep climbs with more gears", "Grab the path with high tread tires",
     "Highest rated Mountain bike", 237, 0.2308, 81, 1000),
    ("Mountain", 2, "WeBike", "HikeBike 1", "Hike Bike",
     "Picture of a rider on a steep trail", "Tackle steep climbs with more gears",
     "Grab the path with high tread tires", 115, 0.1120, 80, 1000),
    ("Mountain", 3, "LiteCycle", "Trail Blazing 1", "LiteTrail Pro",
     "Picture of a rider on a steep trail", "Mountains are no longer difficult",
     "Have an adventure on a carbon bike!", 69, 0.0672, 79, 1000),
    ("Mountain", 4, "Bike Bros", "TerraTech", "TERRAMAX",
     "Picture of a rider on a steep trail", "Go anywhere on a tough carbon bike",
     "Local sales & service", 46, 0.0448, 78, 1000),
    ("Mountain", 5, "Bike Bros", "Mountainability", "TERRAMAX",
     "Picture of a rider on a steep trail", "Have an adventure on a carbon bike!",
     "Tackle steep climbs with more gears", 33, 0.0321, 77, 1000),
    ("Speed", 1, "BB LLC", "Unleash lil pap", "Blu Tube Ballz",
     "Ride a wind-cheater", "Roll fast with racing tires",
     "Elite look - a ride of distinction", 200, 0.2237, 76, 1000),
    ("Speed", 2, "SpaceBikes", "The Armstrong 1", "The Armstrong",
     "Highest rated Speed bike", "Picture of road race",
     "Fast and furious", 98, 0.1096, 77, 1000),
    ("Speed", 3, "LiteCycle", "Speed of Lite 1", "LiteSpeed Pro+",
     "Picture of road race", "Ride a wind-cheater",
     "Elite look - a ride of distinction", 61, 0.0682, 77, 1000),
    ("Speed", 4, "MILC Bikes", "Skim MILC MKII", "Skim MILC MKII",
     "Picture of road race", "Ride a wind-cheater",
     "Elite look - a ride of distinction", 43, 0.0481, 75, 1000),
    ("Speed", 5, "WeBike", "Swift Bike", "Swift Bike",
     "Picture of road race", "Enjoy your ride - carbon fiber light",
     "Carbon fiber quality at a great price", 29, 0.0324, 70, 1000),
    ("Speed", 6, "LiteCycle", "Speed of Lite 2", "LiteSpeed+",
     "Picture of road race", "Enjoy your ride - carbon fiber light",
     "Ride a wind-cheater", 23, 0.0257, 73, 1000),
    ("Speed", 7, "Bike Bros", "AndStill", "MACH I.I",
     "Picture of road race", "Ride a wind-cheater",
     "Elite look - a ride of distinction", 23, 0.0257, 78, 1000),
    ("Speed", 8, "Bike Bros", "Excluspeedity", "MACH I.I",
     "Picture of road race", "Ride a wind-cheater",
     "Fast and furious", 23, 0.0257, 72, 1000),
    ("Speed", 9, "Spoke'd Up", "Spoke'd Speed", "Spoke'd Speed",
     "Fast and furious", "Ride a wind-cheater",
     "Roll fast with racing tires", 21, 0.0235, 57, 1000),
]

WEB = {
    "BB LLC": 569, "SpaceBikes": 491, "LiteCycle": 435, "MILC Bikes": 435,
    "Spoke'd Up": 138, "WeBike": 120, "Bike Bros": 0,
}

print("CTR RECONCILIATION (clicks / searches)")
for seg, pos, co, ad, br, a, b, c, clicks, ctr, aj, inv in SEM:
    expected = clicks / SEARCHES[seg]
    ok = abs(expected - ctr) < 0.00015
    if not ok:
        print(f"  MISMATCH {seg} pos {pos} {ad}: {clicks}/{SEARCHES[seg]} = {expected:.4f} vs {ctr}")
    assert ok
print("  all 18 CTRs = clicks / segment searches\n")

clicks_by = {}
for row in SEM:
    clicks_by[row[2]] = clicks_by.get(row[2], 0) + row[8]
print("COMPANY ORGANIC CLICKS")
for co, n in sorted(clicks_by.items(), key=lambda x: -x[1]):
    web = WEB[co]
    conv = (web / n) if n else 0
    print(f"  {co:<12} clicks {n:>4}  web {web:>4}  web/click {conv:4.2f}")
print(f"\n  WeBike {clicks_by['WeBike']} vs BB LLC {clicks_by['BB LLC']} (chart-read was ~150 vs ~440)")
print(f"  LiteCycle {clicks_by['LiteCycle']} clicks -> {WEB['LiteCycle']} web")
print(f"  WeBike    {clicks_by['WeBike']} clicks -> {WEB['WeBike']} web")
print("  nearly identical traffic, 3.6x web-sales gap -> organic is NOT the web problem\n")

wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
bad = PatternFill("solid", fgColor="FFC7CE")
good = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFEB9C")
ours = PatternFill("solid", fgColor="FFF2CC")

if "Q3_Organic_SEM" in wb.sheetnames:
    del wb["Q3_Organic_SEM"]
ws = wb.create_sheet("Q3_Organic_SEM")

ws["A1"] = "Organic SEM Results — Q3 actual"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("WeBike 144 clicks vs BB LLC 437. HikeBike 1 is #2 in Mountain (115 clicks, ad 80). "
            "Swift Bike is #5 in Speed (29 clicks, ad 70) with Rec/price copy. Organic is NOT why "
            "web sales are last: LiteCycle has 153 clicks and 435 web units; we have 144 and 120.")
ws["A2"].alignment = Alignment(wrap_text=True)


def block(start, rows, wrap_col=None):
    for i, row in enumerate(rows):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(start + i, c, v)
            if i == 0:
                cell.fill = hdr
                cell.font = bold
            elif wrap_col and c == wrap_col:
                cell.alignment = Alignment(wrap_text=True)
    return start + len(rows) + 2


r = 4
rows = [["Segment", "Searches", "Listings", "Total clicks", "Clicks / searches"]]
for seg in ("Recreation", "Mountain", "Speed"):
    n = sum(1 for x in SEM if x[0] == seg)
    clicks = sum(x[8] for x in SEM if x[0] == seg)
    rows.append([seg, SEARCHES[seg], n, clicks, clicks / SEARCHES[seg]])
r2 = block(r, rows)
for i in range(1, 4):
    ws.cell(r + i, 5).number_format = "0.0%"
r = r2

rows = [["Segment", "Pos", "Company", "Ad", "Brand", "SERP 1", "SERP 2", "SERP 3",
         "Clicks", "CTR", "Ad judgment", "Page $"]]
for row in SEM:
    rows.append(list(row))
r2 = block(r, rows)
for i, row in enumerate(SEM, start=1):
    ws.cell(r + i, 10).number_format = "0.00%"
    if row[2] == "WeBike":
        for c in range(1, 13):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
r = r2

rows = [["COMPANY TOTALS", "Organic clicks", "Web demand", "Web units per click", "Note"]]
for co, n in sorted(clicks_by.items(), key=lambda x: -x[1]):
    web = WEB[co]
    conv = (web / n) if n else None
    note = ""
    if co == "WeBike":
        note = "6th of 7; 115 Mountain + 29 Speed + 0 Rec"
    elif co == "BB LLC":
        note = "best; #1 Mountain + #1 Speed, no Rec page"
    elif co == "LiteCycle":
        note = "almost our traffic, 3.6x our web sales"
    elif co == "Bike Bros":
        note = "213 clicks, 0 web centre -> organic without ops captures nothing online"
    rows.append([co, n, web, conv if conv is not None else "n/a", note])
r2 = block(r, rows, wrap_col=5)
for i in range(1, len(rows)):
    if isinstance(rows[i][3], float):
        ws.cell(r + i, 4).number_format = "0.00"
    if rows[i][0] == "WeBike":
        for c in range(1, 6):
            ws.cell(r + i, c).fill = ours
            ws.cell(r + i, c).font = bold
        ws.cell(r + i, 2).fill = bad
    if rows[i][0] == "LiteCycle":
        ws.cell(r + i, 4).fill = yellow
r = r2

rows = [["CTR BY POSITION (clicks / searches) — almost identical across segments",
         "Recreation", "Mountain", "Speed"]]
for pos in range(1, 10):
    rec = next((x[9] for x in SEM if x[0] == "Recreation" and x[1] == pos), None)
    mtn = next((x[9] for x in SEM if x[0] == "Mountain" and x[1] == pos), None)
    spd = next((x[9] for x in SEM if x[0] == "Speed" and x[1] == pos), None)
    rows.append([f"Position {pos}", rec, mtn, spd])
rows.append(["Read", None, None, "Position 1 always takes ~22-23% of searches, position 2 ~11%, "
             "position 3 ~6.7%. CTR is a function of RANK, not ad judgment. Ad judgment sets rank "
             "in Rec and Mountain (monotonic) but NOT in Speed (AndStill 78 sits at #7)."])
r2 = block(r, rows, wrap_col=4)
for i in range(1, 10):
    for c in (2, 3, 4):
        ws.cell(r + i, c).number_format = "0.00%"
r = r2

rows = [
    ["SERP MECHANICS", "Finding"],
    ["Web page investment", "Every listing is $1,000. No one spent more. Page $ is not a Q3 lever."],
    ["Brand name on SERP", "HikeBike 1 locked rank 1 = Mention brand name, but the snippet is picture / "
     "gears / tires (locked ranks 2-4). Brand fills the Ad/Brand columns; it does NOT consume a "
     "content slot. SERP shows the next three benefits."],
    ["Picture-first underperforms", "Rec #1 and Mountain #1 lead with a CLAIM (Highest rated / gears). "
     "Our Mountain page leads with a picture and sits #2. Speed positions 3-8 almost all lead with "
     "the road-race picture; positions 1-2 do not."],
    ["Two pages in one segment cannibalize", "Bike Bros 2 Mountain ads = 79 clicks vs our one page at "
     "115. Their 2 Speed ads both stall at 23 clicks. Do not add a second page."],
]
r2 = block(r, rows, wrap_col=2)
for i in range(1, 5):
    ws.cell(r + i, 2).alignment = Alignment(wrap_text=True)
r = r2

rows = [
    ["COPY vs WINNERS", "SERP 1", "SERP 2", "SERP 3", "Ad j", "Pos", "Clicks"],
    ["BB LLC Mountain (best)", "Tackle steep climbs with more gears",
     "Grab the path with high tread tires", "Highest rated Mountain bike", 81, 1, 237],
    ["WeBike HikeBike 1", "Picture of a rider on a steep trail",
     "Tackle steep climbs with more gears", "Grab the path with high tread tires", 80, 2, 115],
    ["BB LLC Speed (best)", "Ride a wind-cheater", "Roll fast with racing tires",
     "Elite look - a ride of distinction", 76, 1, 200],
    ["WeBike Swift Bike", "Picture of road race", "Enjoy your ride - carbon fiber light",
     "Carbon fiber quality at a great price", 70, 5, 29],
    ["Read", None, None, None, None, None,
     "HikeBike 1 is one point and one picture away from the #1 recipe; Hike Bike IS tied for "
     "highest-rated Mountain so we can claim it. Swift Bike SERP is Rec-coded (enjoy your ride) "
     "plus HikeBike 1's price claim. Speed is the least price-sensitive segment. This CONFIRMS "
     "the Q3_Ad_Judgment clone hypothesis with the actual copy."],
]
r2 = block(r, rows, wrap_col=7)
ws.cell(r + 2, 1).fill = ours
ws.cell(r + 4, 1).fill = ours
ws.cell(r + 4, 5).fill = bad
ws.cell(r + 5, 7).alignment = Alignment(wrap_text=True)
ws.cell(r + 5, 7).fill = yellow
r = r2

rows = [
    ["Q4 SEM ACTIONS (not locked)", "Decision"],
    ["Organic is not the web problem", "144 vs LiteCycle 153 clicks, 120 vs 435 web units. Rebuild "
     "web staff 3->7 and fund all 4 tactics. That is the +315 unit / +$189,800 case."],
    ["Swift Bike ad", "REDESIGN SERP around Speed cues. Target top 3 after brand: wind-cheater / "
     "racing tires / elite look (or Highest rated Speed bike once design hits 77). Drop 'great "
     "price' and 'enjoy your ride'. Already required for Marketing Effectiveness 0.7675."],
    ["HikeBike 1", "KEEP as the default. 80 vs 81, position 2 of 5. Optional cheap tweak if we "
     "are already in Ad Design: rank gears / tires / Highest rated Mountain bike so the snippet "
     "matches BB LLC. Do not pay a full redesign fee for one point."],
    ["Recreation page", "Do not add. No Rec brand; Rec winners all lead with comfort-seat copy we "
     "cannot truthfully run on Hike Bike."],
    ["Second page in a segment", "Do not. Bike Bros split themselves and finished last."],
    ["Page investment", "Leave at $1,000 until a rival demonstrates that spending more moves rank."],
]
block(r, rows, wrap_col=2)
for i in range(1, 7):
    ws.cell(r + i, 2).fill = yellow
    ws.cell(r + i, 2).alignment = Alignment(wrap_text=True)

ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 44
ws.column_dimensions["C"].width = 28
ws.column_dimensions["D"].width = 28
ws.column_dimensions["E"].width = 22
ws.column_dimensions["F"].width = 42
ws.column_dimensions["G"].width = 48
ws.column_dimensions["H"].width = 42
for col in "IJKL":
    ws.column_dimensions[col].width = 14
ws.row_dimensions[2].height = 48

wb.save(DST)
print("Added Q3_Organic_SEM")
