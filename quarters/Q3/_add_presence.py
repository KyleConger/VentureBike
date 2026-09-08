"""Add Q3 competitor presence by city (confirmed store/web footprints) + corrected city strategy."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DST = "quarters/Q3/Q3Data.xlsx"

FIRMS = ["Bike Bros", "LiteCycle", "WeBike", "BB LLC", "Spoke'd Up", "SpaceBikes", "MILC Bikes"]
# confirmed from "Competitors in City" report
PRESENCE = {
    "New York City":  {"BB LLC", "Spoke'd Up", "SpaceBikes"},
    "Rio de Janeiro": {"WeBike"},
    "Amsterdam":      set(FIRMS),
    "Bangalore":      {"LiteCycle"},
}
WEB_CENTRE = {"LiteCycle", "WeBike", "BB LLC", "Spoke'd Up", "SpaceBikes", "MILC Bikes"}

# Q3 demand by firm x city (from Q3_City_Demand, already reconciled)
DEM = {
    "New York City":  {"Bike Bros": 0, "LiteCycle": 143, "WeBike": 38, "BB LLC": 752,
                       "Spoke'd Up": 412, "SpaceBikes": 574, "MILC Bikes": 120},
    "Rio de Janeiro": {"Bike Bros": 0, "LiteCycle": 106, "WeBike": 270, "BB LLC": 139,
                       "Spoke'd Up": 32, "SpaceBikes": 120, "MILC Bikes": 107},
    "Amsterdam":      {"Bike Bros": 490, "LiteCycle": 365, "WeBike": 295, "BB LLC": 493,
                       "Spoke'd Up": 380, "SpaceBikes": 490, "MILC Bikes": 472},
    "Bangalore":      {"Bike Bros": 0, "LiteCycle": 370, "WeBike": 24, "BB LLC": 123,
                       "Spoke'd Up": 30, "SpaceBikes": 114, "MILC Bikes": 100},
}
CITIES = ["Amsterdam", "New York City", "Rio de Janeiro", "Bangalore"]

print("RECONCILIATION")
for city in CITIES:
    print(f"  {city:<16}{sum(DEM[city].values()):>6}")
print(f"  {'TOTAL':<16}{sum(sum(d.values()) for d in DEM.values()):>6}  (report says 6559)")
assert sum(sum(d.values()) for d in DEM.values()) == 6559

print("\nSTORE COUNTS (confirmed)")
for f in FIRMS:
    n = sum(1 for c in CITIES if f in PRESENCE[c])
    print(f"  {f:<12} {n} store(s)  web centre: {'yes' if f in WEB_CENTRE else 'NO'}")

print("\nSTORE FIRMS vs WEB-ONLY FIRMS, capture by city")
for city in CITIES:
    s = sum(v for f, v in DEM[city].items() if f in PRESENCE[city])
    w = sum(v for f, v in DEM[city].items() if f not in PRESENCE[city])
    n = len(PRESENCE[city])
    print(f"  {city:<16} store firms {n}: {s:>5} ({s/(s+w):5.1%})   web-only: {w:>4} ({w/(s+w):5.1%})"
          f"   avg per store firm {s/n:>5.0f}")

wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
bad = PatternFill("solid", fgColor="FFC7CE")
good = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFEB9C")
ours = PatternFill("solid", fgColor="FFF2CC")

if "Q3_Presence" in wb.sheetnames:
    del wb["Q3_Presence"]
ws = wb.create_sheet("Q3_Presence")

ws["A1"] = "Competitors in City — CONFIRMED footprints, Q3"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("This report confirms every store location previously INFERRED from demand patterns. "
            "One correction: MILC Bikes has only ONE store (Amsterdam), not two.")
ws["A2"].alignment = Alignment(wrap_text=True)


def block(start, rows, wrap_col=None):
    for i, row in enumerate(rows):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(start + i, c, v)
            if i == 0:
                cell.fill = hdr
                cell.font = bold
            elif wrap_col and c == wrap_col:
                cell.alignment = Alignment(wrap_text=True)
    return start + len(rows) + 2


# presence matrix
r = 4
rows = [["Location"] + FIRMS + ["Store firms present"]]
for city in CITIES:
    rows.append([city] + ["YES" if f in PRESENCE[city] else "-" for f in FIRMS] + [len(PRESENCE[city])])
rows.append(["Web Sales Centre"] + ["YES" if f in WEB_CENTRE else "-" for f in FIRMS] + [len(WEB_CENTRE)])
r2 = block(r, rows)
for i in range(1, len(rows)):
    for c in range(2, 9):
        if ws.cell(r + i, c).value == "YES":
            ws.cell(r + i, c).fill = ours if FIRMS[c - 2] == "WeBike" else good
r = r2

rows = [["CONFIRMED FOOTPRINT", "Stores", "Cities", "Web centre", "Total demand"]]
for f in sorted(FIRMS, key=lambda x: -sum(DEM[c][x] for c in CITIES)):
    cities = [c for c in CITIES if f in PRESENCE[c]]
    rows.append([f, len(cities), ", ".join(cities), "yes" if f in WEB_CENTRE else "NO",
                 sum(DEM[c][f] for c in CITIES)])
r2 = block(r, rows)
for i in range(1, len(rows)):
    if rows[i][0] == "WeBike":
        for c in range(1, 6):
            ws.cell(r + i, c).fill = ours
    if rows[i][3] == "NO":
        ws.cell(r + i, 4).fill = bad
r = r2

# the structural finding
rows = [["MARKET STRUCTURE — a store never captures the whole city", "Store firms",
         "Store-firm demand", "Store capture %", "Web-only demand", "Web-only %", "Avg per store firm"]]
for city in CITIES:
    s = sum(v for f, v in DEM[city].items() if f in PRESENCE[city])
    w = sum(v for f, v in DEM[city].items() if f not in PRESENCE[city])
    n = len(PRESENCE[city])
    rows.append([city, n, s, s / (s + w), w, w / (s + w), s / n])
r2 = block(r, rows)
for i in range(1, 5):
    ws.cell(r + i, 4).number_format = "0.0%"
    ws.cell(r + i, 6).number_format = "0.0%"
    ws.cell(r + i, 7).number_format = "0"
ws.cell(r + 2, 7).fill = good   # NYC avg per store firm
ws.cell(r + 1, 7).fill = bad    # Amsterdam avg per store firm
r = r2

rows = [
    ["CORRECTION — Amsterdam is the MOST contested city, not an open one", "Note"],
    ["Earlier read", "'No rival exceeds 16.5%, so Amsterdam is genuinely open.'"],
    ["What the presence report shows", "ALL SEVEN firms hold an Amsterdam store. Nobody exceeds 16.5% "
     "precisely BECAUSE all seven are there."],
    ["Avg demand per store firm - Amsterdam", 426],
    ["Avg demand per store firm - New York City", 579],
    ["Implication", "A store in New York is worth MORE than deeper investment in Amsterdam. Amsterdam is "
     "2,985 units split 7 ways; New York is 2,039 units split 3 ways. Gaining share in Amsterdam means "
     "taking it from six store-equipped rivals."],
]
r2 = block(r, rows, wrap_col=2)
ws.cell(r + 1, 2).fill = bad
ws.cell(r + 5, 2).fill = good
r = r2

rows = [
    ["RIO — we under-exploit our monopoly", "WeBike in Rio", "LiteCycle in Bangalore"],
    ["Only firm with a store in the city", "yes", "yes"],
    ["City demand", 774, 761],
    ["Our / their units", 270, 370],
    ["Share as sole store holder", 0.349, 0.486],
    ["Gap", "=C{0}-B{0}".format(r + 4), "13.7 pts of share we leave on the table"],
    ["Web-only rivals take", 504, 391],
    ["Web-only rivals' capture %", "=B{0}/B{1}".format(r + 6, r + 2), "=C{0}/C{1}".format(r + 6, r + 2)],
    ["Read", None, "Same structural position, nearly identical city size - LiteCycle extracts 13.7 more "
     "points than we do. Rio opened only in Q3 and we stocked out, so some of this is timing. But 65% of "
     "OUR exclusive city still goes to rivals who have no store there at all."],
]
r2 = block(r, rows)
ws.cell(r + 4, 2).number_format = "0.0%"
ws.cell(r + 4, 3).number_format = "0.0%"
ws.cell(r + 7, 2).number_format = "0.0%"
ws.cell(r + 7, 3).number_format = "0.0%"
ws.cell(r + 4, 2).fill = bad
ws.cell(r + 4, 3).fill = good
ws.cell(r + 8, 3).alignment = Alignment(wrap_text=True)
ws.cell(r + 8, 3).fill = yellow
r = r2

rows = [
    ["WHY THE WEB CASE IS NOW STRONGER", "Note"],
    ["Evidence", "BB LLC pulls 139 units out of Rio with NO store there - 51% of what we get WITH a store. "
     "SpaceBikes 120, MILC 107, LiteCycle 106. Web-only rivals took 504 of Rio's 774 units."],
    ["Our mirror position", "We pull only 38 from New York and 24 from Bangalore by web. Rivals do 3-5x "
     "better at the same trick."],
    ["Conclusion", "Web reach is not a consolation prize for cities we skip - it is how the strong firms "
     "raid cities they never entered. Our 19.1% web mix is costing us in all four cities, including the "
     "two where we have stores."],
]
r2 = block(r, rows, wrap_col=2)
ws.cell(r + 3, 2).fill = yellow
r = r2

rows = [
    ["CITY #3 — updated read", "New York City", "Bangalore"],
    ["City demand", 2039, 761],
    ["Store firms already present", 3, 1],
    ["Avg demand per store firm today", 579, 370],
    ["Weakest incumbent store firm", "Spoke'd Up 412", "n/a"],
    ["Our current web-only demand there", 38, 24],
    ["Plausible units with a store", "412 - 574 (match Spoke'd Up or SpaceBikes)", "~250 (split LiteCycle's monopoly)"],
    ["Verdict", "STILL THE BEST STORE TARGET. 2.7x the demand, and only 3 rivals hold stores "
     "vs Amsterdam's 7. Enter via SPEED (880 units, no dominant firm) - BB LLC holds 82.8% of NYC Mountain.",
     "Smallest prize. LiteCycle's 48.6% monopoly is the industry's most concentrated position."],
]
r2 = block(r, rows)
for c in (2, 3):
    ws.cell(r + 7, c).alignment = Alignment(wrap_text=True)
ws.cell(r + 7, 2).fill = good
r = r2

rows = [
    ["Q4 SEQUENCING — reconciling depth vs breadth", "Decision"],
    ["1st dollar: web rebuild", "Cheapest demand, reaches all 4 cities, pays back in-quarter (~+$89,600). "
     "Rivals prove web raids cities you never enter."],
    ["2nd dollar: Rio depth", "We are the ONLY store in Rio and extract just 34.9% vs LiteCycle's 48.6% in "
     "the same position. Closing that gap is ~106 units with no new lease."],
    ["3rd dollar: New York store", "Biggest single prize (2,039 units, only 3 store rivals). Enter through "
     "Speed once Swift Bike is fixed to 77."],
    ["NOT a priority", "Deepening Amsterdam against 6 store-equipped rivals; opening Bangalore against "
     "LiteCycle's entrenched 48.6%."],
]
block(r, rows, wrap_col=2)
for i in range(1, len(rows)):
    ws.cell(r + i, 2).fill = yellow

ws.column_dimensions["A"].width = 48
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 46
for col in "DEFGH":
    ws.column_dimensions[col].width = 15

wb.save(DST)
print("\nAdded Q3_Presence")
