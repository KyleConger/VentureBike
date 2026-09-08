"""Build quarters/Q3/Q3Data.xlsx from Q2Data + Q3 actual results (BSC, market, financials)."""
import shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

SRC = "quarters/Q2/Q2Data.xlsx"
DST = "quarters/Q3/Q3Data.xlsx"
shutil.copyfile(SRC, DST)

wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
bad = PatternFill("solid", fgColor="FFC7CE")
good = PatternFill("solid", fgColor="C6EFCE")
yellow = PatternFill("solid", fgColor="FFEB9C")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = bold
    if note:
        ws["A2"] = note


def table(ws, start, rows, widths=None):
    r = start
    for row in rows:
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
        r += 1
    return r


# ---------------- Q3_BSC ----------------
ws = wb.create_sheet("Q3_BSC")
title(ws, "Balanced Scorecard — Q3 actual (WeBike vs industry)",
      "Total Performance = product of all indicators. WeBike = LAST place (0.411 = industry minimum).")
rows = [
    ["Indicator", "Min", "Max", "Average", "WeBike", "vs Avg", "Read"],
    ["Total Performance", 0.411, 21.531, 5.897, 0.411, "=E4-D4", "LAST — dragged by Market Perf + Asset Mgmt"],
    ["Financial Performance", 5.319, 37.687, 19.006, 5.319, "=E5-D5", "LAST — operating loss"],
    ["Market Performance", 0.085, 0.345, 0.185, 0.098, "=E6-D6", "Near last — 33% stock-outs"],
    ["Marketing Effectiveness", 0.660, 0.765, 0.743, 0.738, "=E7-D7", "At par — ads working"],
    ["Investment in Future", 3.466, 7.401, 5.386, 6.795, "=E8-D8", "STRENGTH — 2nd best"],
    ["Wealth", 0.668, 1.017, 0.831, 0.668, "=E9-D9", "LAST — cumulative losses"],
    ["Human Resource Mgmt", 0.691, 0.767, 0.724, 0.713, "=E10-D10", "Below avg — pay fell behind"],
    ["Asset Management", 0.353, 0.920, 0.542, 0.353, "=E11-D11", "LAST — $1.01M idle cash"],
    ["Manufacturing Productivity", 0.722, 0.995, 0.899, 0.938, "=E12-D12", "STRENGTH — but only by under-scheduling"],
    ["Financial Risk", 1.000, 1.000, 1.000, 1.000, "=E13-D13", "No debt"],
]
table(ws, 3, rows)
for c in range(1, 8):
    ws.cell(3, c).fill = hdr
    ws.cell(3, c).font = bold
for r in (4, 5, 9, 11):
    ws.cell(r, 5).fill = bad
for r in (8, 12):
    ws.cell(r, 5).fill = good

table(ws, 15, [
    ["Cumulative BSC (Q2+Q3)", None, None, None, None],
    ["Indicator", "Min", "Max", "Average", "WeBike"],
    ["Cumulative Total Performance", 0.033, 6.540, 2.466, 0.033],
    ["Cumulative Financial Performance", 1.203, 24.032, 12.327, 1.203],
    ["Cumulative Market Performance", 0.091, 0.264, 0.168, 0.091],
    ["Cumulative Marketing Effectiveness", 0.576, 0.754, 0.701, 0.576],
    ["Cumulative Investment in Future", 3.466, 7.401, 5.386, 6.795],
    ["Cumulative Wealth", 0.668, 1.017, 0.831, 0.668],
    ["Cumulative HR Management", 0.666, 0.750, 0.709, 0.733],
    ["Cumulative Asset Management", 0.280, 0.710, 0.444, 0.280],
    ["Cumulative Manufacturing Productivity", 0.566, 0.969, 0.852, 0.566],
    ["Cumulative Financial Risk", 1.000, 1.000, 1.000, 1.000],
])
ws["A16"].font = bold

table(ws, 29, [
    ["Sensitivity — what moves Total Performance", None, None],
    ["Scenario", "Total Performance", "Note"],
    ["Q3 actual", 0.411, "baseline"],
    ["Serve 100% of demand (MP 0.098 -> 0.150)", "=0.411/0.098*0.15", "no stock-outs, same share"],
    ["+ Asset turnover 0.353 -> 0.55", "=B32/0.353*0.55", "deploy idle cash into revenue"],
    ["+ HR 0.713 -> 0.80", "=B33/0.713*0.80", "raise pay above new industry avg"],
])
ws["A29"].font = bold

# ---------------- Q3_Market ----------------
ws = wb.create_sheet("Q3_Market")
title(ws, "Q3 Market Demand & Share (all competitors)",
      "Mountain is our strong segment (#2 of 7). Speed is crowded and we are weak there.")
table(ws, 3, [
    ["Company", "Recreation", "Mountain", "Speed", "Total Demand",
     "Rec %", "Mtn %", "Speed %", "Total %"],
    ["Bike Bros", 138, 172, 180, 490, 0.0670, 0.1061, 0.0625, 0.0747],
    ["LiteCycle", 96, 340, 548, 984, 0.0466, 0.2097, 0.1904, 0.1500],
    ["WeBike", 54, 358, 215, 627, 0.0262, 0.2209, 0.0747, 0.0956],
    ["BB LLC", 99, 751, 657, 1507, 0.0481, 0.4633, 0.2283, 0.2298],
    ["Spoke'd Up", 492, 0, 362, 854, 0.2388, 0.0000, 0.1258, 0.1302],
    ["SpaceBikes", 710, 0, 588, 1298, 0.3447, 0.0000, 0.2043, 0.1979],
    ["MILC Bikes", 471, 0, 328, 799, 0.2286, 0.0000, 0.1140, 0.1218],
    ["TOTAL", 2060, 1621, 2878, 6559, None, None, None, None],
])
for c in range(1, 10):
    ws.cell(3, c).fill = hdr
    ws.cell(3, c).font = bold
for c in range(1, 10):
    ws.cell(6, c).fill = yellow  # WeBike row

table(ws, 14, [
    ["Segment read", "Detail"],
    ["Mountain", "Only 4 of 7 firms compete. We are #2 at 22.1% behind BB LLC 46.3%. Smallest segment (1,621) but least crowded."],
    ["Speed", "Largest segment (2,878) and all 7 firms compete. We are 6th at 7.5%. Swift Bike is not winning."],
    ["Recreation", "Not targeted. 3 firms own ~81% of it."],
    ["BB LLC", "Industry leader: 23% total share, dominates Mountain AND #1 Speed. Main rival in our primary."],
])
ws["A14"].font = bold

table(ws, 21, [
    ["Demand after ill will & stock-outs (WeBike)", None, None, None],
    ["Brand", "Generated Demand", "Net Demand", "Lost to Stock-out", "Units Sold"],
    ["Hike Bike", 412, 412, 134, 278],
    ["Swift Bike", 215, 215, 70, 145],
    ["Total", 627, 627, 204, 423],
    ["Percent of demand served", "=E25/C25", None, None, "67%"],
    ["Q4 ill-will penalty (half of lost %)", "=D25/C25/2", "Q4 generated demand cut by this", None, None],
])
ws["A21"].font = bold

# ---------------- Q3_Financials ----------------
ws = wb.create_sheet("Q3_Financials")
title(ws, "Income Statement / Cash Flow / Balance Sheet — Q1-Q3 actual")
table(ws, 3, [
    ["Income Statement", "Q1", "Q2", "Q3"],
    ["Revenues", 0, 266175, 589720],
    ["- Cost of Goods Sold", 0, 100928, 272340],
    ["= Gross Profit", 0, 165247, 317380],
    ["Research and Development", 30000, 0, 30000],
    ["+ Advertising", 0, 81780, 119165],
    ["+ Internet Marketing", 0, 1000, 2000],
    ["+ Sales Force Expense", 0, 70010, 93815],
    ["+ Store and Web Center Expense", 286000, 194000, 121000],
    ["+ Web Sales Productivity", 0, 18000, 6000],
    ["+ Marketing Research", 0, 20000, 20000],
    ["+ Shipping", 0, 3537, 6830],
    ["+ Excess Capacity Cost", 0, 148652, 0],
    ["+ Depreciation", 0, 30000, 30000],
    ["= Total Expenses", 316000, 566978, 428810],
    ["Operating Profit", -316000, -401731, -111431],
    ["= Net Income", -316000, -401731, -111431],
    ["Earnings per Share", -21, -20, -4],
])
for c in range(1, 5):
    ws.cell(3, c).fill = hdr
    ws.cell(3, c).font = bold

table(ws, 24, [
    ["Cash Flow", "Q1", "Q2", "Q3"],
    ["Beginning Cash", 0, 464000, 592269],
    ["Net Operating Cash Flow", -316000, -371731, -81431],
    ["Investing (fixed capacity)", 720000, 0, 0],
    ["Financing (stock issued)", 1500000, 500000, 500000],
    ["Ending Cash", 464000, 592269, 1010838],
])
ws["A24"].font = bold

table(ws, 32, [
    ["Balance Sheet", "Q1", "Q2", "Q3"],
    ["Cash", 464000, 592269, 1010838],
    ["3-Month CD", 0, 0, 0],
    ["Net Fixed Assets", 720000, 690000, 660000],
    ["= Total Assets", 1184000, 1282269, 1670838],
    ["Common Stock", 1500000, 2000000, 2500000],
    ["Retained Earnings", -316000, -717731, -829162],
    ["Shares issued", 15000, 20000, 25000],
])
ws["A32"].font = bold

table(ws, 42, [
    ["Industry ratios (Q3)", "Lowest", "Highest", "Average", "WeBike"],
    ["Fixed Assets Turnover", 0.63, 2.14, 1.22, 0.89],
    ["Total Assets Turnover", 0.35, 0.92, 0.54, 0.35],
    ["Gross Profit Margin", 0.5382, 0.6547, 0.5824, 0.5382],
    ["Net Profit Margin", -0.41, 0.2823, 0.009, -0.189],
    ["Return on Assets", -0.1493, 0.2132, 0.0368, -0.0667],
    ["Revenues", 589720, 2198310, 1170036, 589720],
    ["Gross Profit", 317380, 1204529, 683601, 317380],
    ["Net Income", -297348, 510317, 103771, -111431],
])
ws["A42"].font = bold

# ---------------- Q3_Unit_Economics ----------------
ws = wb.create_sheet("Q3_Unit_Economics")
title(ws, "Q3 realized unit economics + cost of the stock-out",
      "Actual prices/costs from the Sales report. Contribution per unit is strong; we simply did not build enough.")
table(ws, 3, [
    ["Brand", "Units Sold", "Revenue", "COGS", "Price/unit", "Cost/unit", "Contribution/unit", "Lost units", "Lost contribution"],
    ["Hike Bike", 278, 379470, 171741, "=C4/B4", "=D4/B4", "=E4-F4", 134, "=G4*H4"],
    ["Swift Bike", 145, 210250, 100599, "=C5/B5", "=D5/B5", "=E5-F5", 70, "=G5*H5"],
    ["Total", "=SUM(B4:B5)", "=SUM(C4:C5)", "=SUM(D4:D5)", None, None, None, "=SUM(H4:H5)", "=SUM(I4:I5)"],
])
for c in range(1, 10):
    ws.cell(3, c).fill = hdr
    ws.cell(3, c).font = bold

table(ws, 9, [
    ["Counterfactual — if every unit of demand had been built", None],
    ["Q3 operating profit (actual)", -111431],
    ["+ Gross margin lost to stock-outs", "=I6"],
    ["= Operating profit if fully served", "=B10+B11"],
    ["Note", "Ignores extra shipping/labor; higher volume would also LOWER unit cost (cost curve)."],
])
ws["A9"].font = bold
ws["B12"].fill = good

table(ws, 16, [
    ["Capacity back-solve — what actually happened", None],
    ["Units produced", 423],
    ["Overtime as % of operating capacity used", 0.13],
    ["Implied normal-hours units", "=B17/(1+B18)"],
    ["Implied units/day (65 days)", "=B19/65"],
    ["Production worker productivity (actual)", 0.72],
    ["Implied SCHEDULED operating capacity", "=B20/B21"],
    ["Fixed capacity available", 24],
    ["Unused fixed capacity per day", "=B23-B22"],
    ["Read", "We owned 24/day and only scheduled ~8/day. The stock-out was a scheduling error, not a printer shortage."],
])
ws["A16"].font = bold
ws["B22"].fill = bad
ws["B24"].fill = bad

# ---------------- Q4_Planner ----------------
ws = wb.create_sheet("Q4_Planner")
title(ws, "Q4 planning — operating capacity sizing (yellow = decision)",
      "Rule: OC/day = forecast demand / 65 * (1 + (1 - productivity)). Never schedule below forecast again.")
table(ws, 3, [
    ["Input", "Value", "Note"],
    ["Q3 generated demand", 627, "actual"],
    ["Ill-will penalty % (half of 32.5% lost)", 0.163, "applies to Q4 generated demand"],
    ["Assumed demand growth before ill will", 1.25, "YELLOW — more people, Rio live, ads maturing"],
    ["Q4 forecast demand (after ill will)", "=B4*B6*(1-B5)", "feed to Demand_Forecast"],
    ["Assumed production worker productivity", 0.72, "YELLOW — raise with better pay"],
    ["Productivity pad", "=1+(1-B8)", None],
    ["Required OC per day", "=B7/65*B9", "schedule AT LEAST this"],
    ["Required OC per quarter", "=B10*65", None],
    ["Fixed capacity per day", 24, "3 printers"],
    ["Headroom (fixed - required)", "=B12-B10", "if negative, order printers THIS quarter"],
    ["Suggested scheduled OC (safety on top)", "=ROUNDUP(B10*1.1,0)", "~10% cushion vs excess-capacity cost"],
])
for c in range(1, 4):
    ws.cell(3, c).fill = hdr
    ws.cell(3, c).font = bold
for r in (6, 8, 16):
    ws.cell(r, 2).fill = yellow

table(ws, 18, [
    ["Other Q4 levers", "Q3 value", "Direction", "Why"],
    ["Sales force productivity", 0.70, "RAISE PAY", "Fell from 85% projected — competitors out-paid us; Q4 shows industry averages"],
    ["Production worker productivity", 0.72, "RAISE PAY", "Same; also feeds capacity pad"],
    ["Common stock issued", 500000, "STOP if cash idle", "More shares lowers Financial Performance (per-share) and Asset Turnover"],
    ["Cash on hand", 1010838, "DEPLOY", "Idle cash is the #1 cause of Asset Mgmt 0.353 (last place)"],
    ["City #3", None, "OPEN ONE", "Revenue growth is the cure for Asset Turnover and unit cost"],
    ["Speed share", 0.0747, "FIX OR REFOCUS", "Speed ad judgment 70 vs Mountain 80; segment is crowded by all 7 firms"],
    ["Gross margin", 0.5382, "RISES WITH VOLUME", "Lowest in industry because volume is lowest"],
])
ws["A18"].font = bold

wb.save(DST)
print(f"Wrote {DST}")
print("New sheets: Q3_BSC, Q3_Market, Q3_Financials, Q3_Unit_Economics, Q4_Planner")
