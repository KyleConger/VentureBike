"""Correct the Q4_Planner scenario read: Rio + full web rebuild EXCEEDS 24/day fixed capacity."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DST = "quarters/Q3/Q3Data.xlsx"
wb = load_workbook(DST)
ws = wb["Q4_Planner"]
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
yellow = PatternFill("solid", fgColor="FFEB9C")
good = PatternFill("solid", fgColor="C6EFCE")
bad = PatternFill("solid", fgColor="FFC7CE")

ws["D53"] = ("Scenario C (Rio fill) fits comfortably at 17.8/day. Scenario D (Rio + full web rebuild) needs "
             "24.3/day and EXCEEDS our 24/day fixed capacity by a hair. Scenario E (plus New York) needs "
             "32.7/day. Capacity - not demand, not product - is now the binding constraint on growth.")
ws["D53"].alignment = Alignment(wrap_text=True)
ws["D53"].fill = bad
ws["D52"].fill = bad
ws["D51"].fill = bad

r = 66
rows = [
    ["CAPACITY IS NOW THE BINDING CONSTRAINT — and idle cash is the cure", "Value", "Note"],
    ["Cash sitting idle", 1010838, "the direct cause of Asset Management 0.353 (LAST place)"],
    ["Fixed capacity today", 24, "3 printers, $720,000 original cost -> ~$240,000 each"],
    ["Max units at 24/day and 74% productivity", "=24*0.74*65", None],
    ["Demand if we fill Rio AND rebuild web", 1171, "exceeds the ceiling by ~17 units"],
    ["A 4th printer would give", 32, "units/day -> max ~1,539 units"],
    ["Approx capital cost", 240000, "24% of idle cash"],
    ["Approx added depreciation", 10000, "per quarter"],
    ["Read", None, "Buying capacity solves TWO problems at once: it lifts the growth ceiling AND converts "
     "idle cash into revenue-producing assets, which is exactly what Asset Management 0.353 is punishing. "
     "Note that owning capacity is NOT what caused Q2's $148,652 excess-capacity charge - that came from "
     "SCHEDULING operating capacity we did not use. Depreciation is the only carrying cost of ownership."],
    ["Decision framing", None, "If we want Rio + web + New York we need roughly 33/day, i.e. 2 more printers. "
     "If we want Rio + web only, 1 more printer gives comfortable slack. Doing nothing caps us at ~1,154 units."],
]
for i, row in enumerate(rows):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r + i, c, v)
        if i == 0:
            cell.fill = hdr
            cell.font = bold
ws.cell(r + 1, 2).fill = bad
ws.cell(r + 5, 2).fill = good
ws.cell(r + 8, 3).alignment = Alignment(wrap_text=True)
ws.cell(r + 8, 3).fill = yellow
ws.cell(r + 9, 3).alignment = Alignment(wrap_text=True)
ws.cell(r + 9, 3).fill = yellow

wb.save(DST)
print("Q4_Planner scenario read corrected; printer/idle-cash block added")
