"""Q4 social media: start Social Networking or wait? Plus revised Q4 cash after the media cut."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DST = "quarters/Q3/Q3Data.xlsx"

CASH = 1_010_838
LOAN = 950_000
FLOOR = 300_000
CAP_Q4 = 1_154
FCST_Q4 = 1_171
CAP_Q5 = round(40 * 0.74 * 65)

print("=" * 98)
print("PART 1 — WHY THIS IS THE OPPOSITE CASE FROM MEDIA INSERTS")
print("=" * 98)
print(f"""  An hour ago I argued for CUTTING media inserts because Q4 output is capped at {CAP_Q4:,} units
  against a forecast of {FCST_Q4:,}, and demand we cannot build only lowers the percent-served term
  in Market Performance. That logic does NOT carry over to social media, and the reason is timing
  rather than any change of principle.

    MEDIA INSERTS   deliver their demand THIS quarter, at full strength.
                    -> harmful when capacity is capped.
    SOCIAL MEDIA    delivers almost NO demand this quarter. The brief is explicit: 'do not expect
                    large increases in demand during your early use of social media.'
                    -> harmless when capacity is capped.

  AND THE RAMP CURVE LINES UP WITH OUR CAPACITY CURVE:
    Q4  capacity {CAP_Q4:>6,} units   social media contributes ~nothing   <- we cannot use demand
    Q5  capacity {CAP_Q5:>6,} units   social media starting to work       <- we NEED demand
    Q6  capacity {CAP_Q5:>6,} units   social media near peak              <- we NEED demand

  Two printers land in Q5 and take us from 24/day to 40/day. Add the New York store (340-570
  units) and Q5 needs several hundred units of NEW demand. A channel that matures on exactly
  that schedule is well timed, not badly timed.""")

print("\n" + "=" * 98)
print("PART 2 — THE DECIDING ARGUMENT: THE CLOCK ONLY RUNS ONCE YOU START")
print("=" * 98)
print("""  From the decision tip, the two sources of delay:
    1. Search engines have to find us and get us cross-listed and referenced.
    2. Our own social media team is inexperienced and has to learn by trial and error.

  And the line that settles it: 'if you decide to wait, you just push back the time it takes for
  search engines to do their work and your social media team to learn.'

  So the ramp is NOT calendar-driven, it is start-date-driven. Waiting a quarter does not let us
  skip the learning curve later — it moves the whole curve a quarter to the right. Every quarter
  we delay is a quarter of maturity we can never recover.

  This is the same reasoning we used to schedule the R&D projects as 1-quarter rapid: when the
  scarce resource is TIME rather than money, buy time.""")

print("\n" + "=" * 98)
print("PART 3 — NO COMPETITOR BENCHMARK EXISTS (verified)")
print("=" * 98)
print("""  I checked all 52 sheets of the Q3 workbook and every competitor dump we hold — Competitor
  Profiles, Media, Organic SEM, Ad Judgment, Web Ops. There is NO social media data anywhere.
  The Q2 internet-marketing background listed social as a tool that 'unlocks over later
  quarters', and this is that unlock.

  MEANING: this is a brand-new decision for EVERY firm, not just us. For once we are not behind.
  Every rival faces the same unknown effectiveness and the same ramp delay starting now.

  It also means we have no spending benchmark. The closest anchors inside the sim are the
  per-activity spend units we HAVE observed:
    organic SEM page      $1,000   (every firm, industry-wide)
    ad design             $6,000   per ad
    Biking media insert   $4,500   each
    web toll-free phone   $9,000   BB LLC
    web page upgrades     $9,000   BB LLC
    web order tracking    $8,000   BB LLC
    web shopping cart     $7,000   BB LLC
  The sim's unit of a meaningful single activity is roughly $5,000-$10,000 a quarter.""")

print("\n" + "=" * 98)
print("PART 4 — WHICH MEDIUM? THE AGENCY SAYS SOCIAL NETWORKING; THE NEEDS DATA AGREES")
print("=" * 98)
NEEDS_IMAGE = [("Status symbol/exclusivity", 121, 116), ("Colorful", 120, 109),
               ("Stylish", 118, 116), ("Feel young at heart", 115, 117),
               ("Fits rider's personality", 113, 116)]
NEEDS_EXPERT = [("Rider expertise required", 123, 126), ("Provides competitive advantage", 117, 124)]
print(f"  Identity and image needs — what Social Networking is good at:")
print(f"    {'Need':<34}{'Mountain':>10}{'Speed':>8}")
for n, m, s in NEEDS_IMAGE:
    print(f"    {n:<34}{m:>10}{s:>8}")
print(f"\n  Expertise and status needs — what a Blog would target:")
for n, m, s in NEEDS_EXPERT:
    print(f"    {n:<34}{m:>10}{s:>8}")
print("""
  Social Networking maps onto a cluster of five needs that all score 109-121 in BOTH our
  segments. That is the same cluster the new bright-colourful-decals R&D project targets, so the
  two reinforce each other. No reason to override the agency.

  MY PREFERENCE FOR THE Q5 ADDITION (pre-registering it now): VIDEO SHARING. Bicycles are an
  inherently visual, action product, and the single most-used benefit in the entire industry is a
  PICTURE — 'Picture of a rider on a steep trail' runs on 4 of 4 Mountain ads and 'Picture of
  road race' on 8 of 8 Speed ads. Video is the natural extension of the one creative device
  every firm already agrees on.
  Q6 candidate: BLOG/NEWSLETTER, which fits 'Rider expertise required' (Mtn 123, Speed 126) —
  one of the highest-scoring needs we currently do nothing about.""")

print("\n" + "=" * 98)
print("PART 5 — HOW MUCH? START SMALL, AND FOR A SPECIFIC REASON")
print("=" * 98)
print("""  The brief says budget drives headcount and tooling, and 'the more you spend, the better your
  potential performance'. But it ALSO says the delay comes from the team's inexperience and that
  the skill 'will come through trial and error'. Those two statements pull apart: a bigger team
  during the ramp is a bigger team that does not yet know what it is doing.

  RECOMMEND $10,000 for Q4. Reasons:
    - It starts the clock, which is the only thing that actually matters this quarter.
    - It sits at the top of the sim's observed per-activity unit ($5,000-$10,000), so it is a
      credible rather than a token effort.
    - It buys the INFORMATION cheaply. Q4 gives us the first real read on whether social media
      does anything for carbon bikes, and we scale in Q5 with evidence instead of hope.
    - It is 1% of the loan and leaves the cash floor untouched.

  THAT LAST POINT IS THE SAME LOGIC THE TEAM ALREADY ACCEPTED ON R&D: we skipped the third R&D
  slot specifically so the decals result would CALIBRATE what an R&D upgrade is worth before
  committing $1M. Identical principle here — buy the measurement first, then size the bet.

  Alternatives:
    $5,000   pure clock-start. Cheapest way to avoid losing a quarter of ramp.
    $20,000  if the team believes trial-and-error learning scales with headcount. Defensible,
             but it doubles a bet on a mechanism the sim openly says is unproven.
    $0       only defensible if the sim ends at Q5. See the caveat below.""")

print("\n" + "=" * 98)
print("PART 6 — REVISED Q4 CASH (the media cut more than pays for this)")
print("=" * 98)
USES = [
    ("R&D — decals + 11-speed, listed", 677_669, ""),
    ("R&D rapid premium (~30% ESTIMATE)", 203_301, "biggest unknown in the plan"),
    ("Operating plan, recurring", 75_000, "was $146,000 — media swings from +$27,000 to -$44,000 under Option C"),
    ("Operating plan, one-time", 67_000, "hiring and setup"),
    ("Swift Bike brand redesign", 30_000, "14-speed + decals => 77"),
    ("Swift Bike ad redesign", 10_000, "AndStill's stack"),
    ("HikeBike 1 ad redesign", 10_000, "BB LLC's stack — this was missing from my earlier table"),
    ("Printers #4 and #5", 480_000, "arrive Q5, when New York opens"),
    ("SOCIAL NETWORKING", 10_000, "NEW — starts the ramp clock"),
]
tot = sum(u[1] for u in USES)
print(f"  {'Use':<40}{'Amount':>12}")
for n, a, note in USES:
    print(f"  {n:<40}{a:>12,}")
print(f"  {'-'*52}")
print(f"  {'TOTAL USES':<40}{tot:>12,}")
print(f"\n  {'Cash on hand':<40}{CASH:>12,}")
print(f"  {'Bank loan':<40}{LOAN:>12,}")
print(f"  {'Total available':<40}{CASH+LOAN:>12,}")
print(f"  {'Less uses':<40}{-tot:>12,}")
print(f"  {'ENDING CASH':<40}{CASH+LOAN-tot:>12,}")
print(f"  {'Hard floor':<40}{FLOOR:>12,}")
print(f"  {'HEADROOM ABOVE FLOOR':<40}{CASH+LOAN-tot-FLOOR:>12,}")
print(f"""
  The media decision funded this outright. Option C takes inserts from $116,000 to $72,000, and
  versus the superseded Option B (+$27,000) that is a ${71_000:,} swing — seven times the social
  media budget. We are not adding spend, we are REALLOCATING it from a channel that would have
  generated demand we cannot build into one that generates demand exactly when we can.""")

print("\n" + "=" * 98)
print("PART 7 — THE HONEST CASE AGAINST, AND WHY IT LOSES")
print("=" * 98)
print(f"""  1. ATTRIBUTION WILL BE IMPOSSIBLE. Q4 already changes R&D, two printers, two ad rebuilds, a
     brand redesign, a price rise, the media mix, web staffing, web tactics, Rio staffing,
     compensation and a $950,000 loan. Adding a twelfth variable means we cannot isolate what
     social media did — which directly undercuts the brief's own instruction to 'monitor the
     effectiveness of your expenditures very carefully'.
     REBUTTAL: real, but it argues for a SMALL STEADY budget we do not fiddle with, not for
     delay. Delay does not improve attribution; it just costs a quarter of ramp.

  2. WE ARE LAST AND SHOULD FIX BASICS FIRST. Total Performance 0.411, last in the industry, on
     stock-outs and idle cash. Social media fixes neither.
     REBUTTAL: also true, and it is exactly why the budget should be $10,000 rather than $50,000.
     At 1% of the loan it does not compete with the capacity and web fixes for resources.

  3. EFFECTIVENESS MAY BE ZERO. The brief openly says social media works for some products and
     has 'little impact' for others.
     REBUTTAL: this is why we buy the small version. And note 'You can stop these activities at
     any time' — this is a cheap OPTION, not a commitment.

  4. THE REAL CAVEAT — HOW MANY QUARTERS REMAIN? This is the same unanswered question that hangs
     over the R&D scheduling decision. If the simulation ends at Q5, a channel that needs 'a few
     quarters' to mature never pays back and the honest answer is $0. If it runs to Q6 or beyond,
     starting now is clearly right. Nobody has told us, and it is the one input that could flip
     this. The asymmetry is what makes me recommend starting anyway: the downside is ${10_000:,},
     the upside is a matured demand channel, and we can stop at any time.""")

print("\n" + "=" * 98)
print("PART 8 — RECOMMENDATION")
print("=" * 98)
print("""  YES — START SOCIAL NETWORKING AT $10,000 FOR Q4.

    - Follow the agency: one medium now, add one per quarter. Q5 Video sharing, Q6 Blog.
    - Expect NOTHING measurable this quarter. Judge it on the Q5 and Q6 trend, not Q4.
    - Do not raise the budget in Q5 unless we see a signal. Do not cut it either — cutting
      restarts the very clock we are paying to start.
    - Record the $10,000 now so we have a clean baseline to measure against.

  The whole case is timing. It costs 1% of the loan, generates no demand in the one quarter where
  demand would hurt us, and matures precisely when two new printers and a New York store leave us
  needing several hundred units of new demand.""")
print("=" * 98)

# ---------------------------------------------------------------- workbook
wb = load_workbook(DST)
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="DDEBF7")
good = PatternFill("solid", fgColor="C6EFCE")
badf = PatternFill("solid", fgColor="FFC7CE")
yel = PatternFill("solid", fgColor="FFEB9C")
thin = Border(*[Side(style="thin", color="D0D0D0")] * 4)

if "Q4_Social_Media" in wb.sheetnames:
    del wb["Q4_Social_Media"]
ws = wb.create_sheet("Q4_Social_Media")


def block(start, rws, wrap=()):
    for i, row in enumerate(rws):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(start + i, c, v)
            cell.border = thin
            if i == 0:
                cell.fill = hdr
                cell.font = bold
            if c in wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    return start + len(rws) + 2


ws["A1"] = "Q4 social media — START Social Networking at $10,000"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("THE CASE IS ENTIRELY ABOUT TIMING. Media inserts deliver demand THIS quarter, which is why we cut "
            "them — Q4 output is capped at 1,154 units against a 1,171 forecast. Social media delivers almost NO "
            "demand this quarter (the brief says so explicitly) and matures in Q5-Q6, exactly when two new printers "
            "take capacity to 1,924 units and a New York store opens. Its ramp curve matches our capacity curve. "
            "And the clock only runs once we start: waiting does not skip the learning curve, it shifts it right.")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:E2")
ws.row_dimensions[2].height = 58

r = 4
rws = [["TIMING — ramp curve vs capacity curve", "Capacity (units)", "Social media state", "Can we use new demand?"]]
for q, cap, st, use in [("Q4", CAP_Q4, "contributes ~nothing", "NO — capped, forecast already 1,171"),
                        ("Q5", CAP_Q5, "starting to work", "YES — 2 printers land, New York opens"),
                        ("Q6", CAP_Q5, "near peak", "YES — needs several hundred new units")]:
    rws.append([q, cap, st, use])
r2 = block(r, rws)
for i in range(1, len(rws)):
    ws.cell(r + i, 2).number_format = "#,##0"
    ws.cell(r + i, 4).fill = badf if i == 1 else good
r = r2

rws = [["WHY NOT WAIT", "Detail"]]
for a, b in [
    ("The clock is start-date driven, not calendar driven", "The tip is explicit: 'if you decide to wait, you just "
     "push back the time it takes for search engines to do their work and your social media team to learn.' Delaying "
     "does not let us skip the ramp later — it moves the whole curve right. Every quarter of delay is a quarter of "
     "maturity we can never recover."),
    ("Two independent sources of delay", "Search engines must find and cross-list us, AND our own team must learn by "
     "trial and error. Neither can be bought later."),
    ("Same logic as the rapid R&D schedule", "When the scarce resource is TIME rather than money, buy time. The team "
     "already accepted that reasoning when scheduling both R&D projects as 1-quarter."),
]:
    rws.append([a, b])
r2 = block(r, rws, wrap=(2,))
r = r2

rws = [["NO COMPETITOR BENCHMARK EXISTS (verified across all 52 sheets)", "Detail"]]
for a, b in [
    ("Zero social media data anywhere", "Checked Competitor Profiles, Media, Organic SEM, Ad Judgment and Web Ops. "
     "Nothing. The Q2 internet-marketing background listed social as a tool that 'unlocks over later quarters' — "
     "this is that unlock."),
    ("So EVERY firm starts from zero", "For once we are not behind. Every rival faces the same unknown effectiveness "
     "and the same ramp delay, starting now."),
    ("Spend anchors we DO have", "Organic SEM page $1,000 (industry-wide) · ad design $6,000 · Biking insert $4,500 · "
     "BB LLC web tactics: toll-free $9,000, page upgrades $9,000, order tracking $8,000, cart $7,000. The sim's unit "
     "for one meaningful activity is roughly $5,000-$10,000 a quarter."),
]:
    rws.append([a, b])
r2 = block(r, rws, wrap=(2,))
r = r2

rws = [["WHICH MEDIUM — needs evidence", "Mountain", "Speed", "Fit"]]
for n, m, s in NEEDS_IMAGE:
    rws.append([n, m, s, "Social Networking (identity/image cluster)"])
for n, m, s in NEEDS_EXPERT:
    rws.append([n, m, s, "Blog / Newsletter — Q6 candidate"])
r2 = block(r, rws, wrap=(4,))
for i in range(1, len(NEEDS_IMAGE) + 1):
    for c in (1, 2, 3, 4):
        ws.cell(r + i, c).fill = good
r = r2
ws.cell(r - 1, 1, "Social Networking maps onto five needs that ALL score 109-121 in BOTH our segments — the same "
                  "cluster the bright-colourful-decals R&D project targets, so the two reinforce each other. No "
                  "reason to override the agency. PRE-REGISTERING Q5: VIDEO SHARING. Bicycles are a visual action "
                  "product and the most-used benefit in the whole industry is a PICTURE ('rider on a steep trail' on "
                  "4 of 4 Mountain ads, 'road race' on 8 of 8 Speed ads). Video extends the one creative device every "
                  "firm already agrees on. Q6: Blog, for 'Rider expertise required' (Mtn 123, Speed 126) — a "
                  "top-scoring need we currently do nothing about.")
ws.cell(r - 1, 1).alignment = Alignment(wrap_text=True)

rws = [["BUDGET OPTIONS", "Amount", "Read"]]
for a, b, c in [
    ("Skip it", 0, "Only defensible if the simulation ends at Q5. A channel needing 'a few quarters' to mature would never pay back."),
    ("Pure clock-start", 5_000, "Cheapest way to avoid losing a quarter of ramp."),
    ("RECOMMENDED", 10_000, "Top of the sim's observed per-activity unit, so credible rather than token. Starts the clock, buys the MEASUREMENT cheaply, and is 1% of the loan. Same principle the team already accepted on R&D: calibrate first, then size the bet."),
    ("Aggressive", 20_000, "Only if the team believes trial-and-error learning scales with headcount. The brief pulls both ways — budget drives headcount, but the delay is attributed to INEXPERIENCE, and a bigger team during the ramp is a bigger team that does not yet know what it is doing."),
]:
    rws.append([a, b, c])
r2 = block(r, rws, wrap=(3,))
for i in range(1, len(rws)):
    ws.cell(r + i, 2).number_format = "$#,##0"
for c in (1, 2, 3):
    ws.cell(r + 3, c).fill = good
    ws.cell(r + 3, c).font = bold
r = r2

rws = [["REVISED Q4 CASH", "Amount", "Note"]]
for n, a, note in USES:
    rws.append([n, a, note])
rws += [["TOTAL USES", tot, ""], ["Cash on hand", CASH, ""], ["Bank loan", LOAN, "7% annual"],
        ["Total available", CASH + LOAN, ""], ["ENDING CASH", CASH + LOAN - tot, ""],
        ["Hard floor", FLOOR, "ending Cash + CD must stay at or above"],
        ["HEADROOM ABOVE FLOOR", CASH + LOAN - tot - FLOOR, ""]]
r2 = block(r, rws, wrap=(3,))
for i in range(1, len(rws)):
    ws.cell(r + i, 2).number_format = "#,##0"
for c in (1, 2, 3):
    ws.cell(r + 9, c).fill = good
    ws.cell(r + 9, c).font = bold
    ws.cell(r + len(rws) - 1, c).fill = yel
    ws.cell(r + len(rws) - 1, c).font = bold
r = r2
ws.cell(r - 1, 1, "THE MEDIA DECISION FUNDED THIS OUTRIGHT. Option C takes inserts from $116,000 to $72,000; against "
                  "the superseded Option B (+$27,000) that is a $71,000 swing, seven times the social media budget. "
                  "We are not adding spend — we are REALLOCATING it out of a channel that would have generated demand "
                  "we cannot build, into one that generates demand exactly when we can. Note the HikeBike 1 ad "
                  "redesign fee was missing from my earlier sources-and-uses table and is now included.")
ws.cell(r - 1, 1).alignment = Alignment(wrap_text=True)

rws = [["THE CASE AGAINST", "Rebuttal"]]
for a, b in [
    ("Attribution will be impossible — Q4 already changes 11 other things",
     "Real. But it argues for a SMALL STEADY budget we do not fiddle with, not for delay. Delay does not improve "
     "attribution; it just costs a quarter of ramp."),
    ("We are last and should fix basics first",
     "Also true, and precisely why the number is $10,000 not $50,000. At 1% of the loan it does not compete with the "
     "capacity and web fixes for resources."),
    ("Effectiveness may be zero for carbon bikes",
     "Which is why we buy the small version. And 'You can stop these activities at any time' — this is a cheap "
     "OPTION, not a commitment."),
    ("HOW MANY QUARTERS REMAIN? — the real caveat",
     "Same unanswered question hanging over the R&D schedule. If the sim ends at Q5, the honest answer is $0. If it "
     "runs to Q6+, starting now is clearly right. The ASYMMETRY is why I still recommend starting: downside $10,000, "
     "upside a matured demand channel, and we can stop any time."),
]:
    rws.append([a, b])
r2 = block(r, rws, wrap=(2,))
ws.cell(r + 4, 2).fill = yel
r = r2

rws = [["DECISION", "Action"]]
for a, b in [
    ("Start Social Networking", "$10,000 for Q4."),
    ("Follow the agency's cadence", "One medium now, add one per quarter. Q5 Video sharing, Q6 Blog/Newsletter."),
    ("Expect nothing this quarter", "Judge it on the Q5 and Q6 trend, not on Q4 results."),
    ("Do not cut it later", "Cutting restarts the very clock we are paying to start. Hold the budget unless we see a signal."),
    ("Record the baseline", "Log the $10,000 now so Q5 has something clean to measure against."),
    ("Status", "ANALYSIS ONLY — majority vote required."),
]:
    rws.append([a, b])
block(r, rws, wrap=(2,))

ws.column_dimensions["A"].width = 50
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 100
ws.column_dimensions["D"].width = 44
ws.column_dimensions["E"].width = 14

wb.save(DST)
print("\nAdded Q4_Social_Media to Q3Data.xlsx")
